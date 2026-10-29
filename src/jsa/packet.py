"""Application packet folders (PRD 04): the job description and the user's own resume and cover letter copies."""

import glob
import shutil
from pathlib import Path
from typing import NamedTuple

from docx.document import Document

from jsa import db
from jsa.docx_pdf import ensure_pdf
from jsa.errors import JsaError
from jsa.naming import (
    cover_letter_file_stem,
    packet_dir_name,
    resume_file_stem,
)
from jsa.placeholders import TITLE, fill_placeholders, holds_placeholder
from jsa.profile import (
    COVER_LETTER_PREFIX,
    Config,
    PacketSources,
    RoleFamily,
    load_config,
    packet_sources,
    with_role_family,
)
from jsa.resume import load_resume
from jsa.tracker import today

JOB_POSTING = "job_posting.md"
REDLINE_EDITS = "redline_edits.json"


def packet_paths(config: Config, job: db.PacketJob) -> tuple[Path, Path]:
    """The packet's folder and the resume copy inside it."""
    directory = config.packets_dir / packet_dir_name(
        job.normalized_company, job.title_slug, job.id, shares_name=job.shares_name
    )
    copy = (
        directory
        / f"{resume_file_stem(config.candidate_name, job.title_slug, job.normalized_company)}.docx"
    )
    return directory, copy


def cover_letter_stem(config: Config, job: db.PacketJob) -> str:
    return cover_letter_file_stem(
        config.candidate_name, job.title_slug, job.normalized_company
    )


class PacketCopy(NamedTuple):
    """A file the packet holds a user-owned copy of, and the profile file it is copied from."""

    path: Path
    source: Path


def packet_copies(
    config: Config, job: db.PacketJob, sources: PacketSources
) -> list[PacketCopy]:
    """The copies of the job's role family's files, the resume copy last; the caller has checked the profile has the role family."""
    directory, resume_copy = packet_paths(config, job)
    family = sources.role_families[job.role_family]
    copies = []
    if family.cover_letter is not None:
        # The copy keeps the source's extension, whatever it is (`.docx`, `.pdf`, ...).
        extension = family.cover_letter.name.removeprefix(COVER_LETTER_PREFIX)
        copies.append(
            PacketCopy(
                directory / f"{cover_letter_stem(config, job)}{extension}",
                family.cover_letter,
            )
        )
    # Last, so the resume copy's existence marks the head as complete (see `ensure_head`).
    copies.append(PacketCopy(resume_copy, family.resume))
    return copies


def _placeholder_document(path: Path) -> Document | None:
    """The `.docx` at `path`; None for any other file, or one that can't be read and so holds no placeholder."""
    if path.suffix.lower() != ".docx":
        return None
    try:
        return load_resume(path)
    except JsaError:
        return None


def _copy_filled(source: Path, copy: Path, job: db.PacketJob) -> None:
    """Copy a profile file into the packet, filling its placeholders; a file with none is copied as it is."""
    document = _placeholder_document(source)
    if document is not None and fill_placeholders(
        document, job.normalized_company, job.title, today()
    ):
        document.save(str(copy))
    else:
        shutil.copyfile(source, copy)


def holds_title_placeholder(family: RoleFamily) -> bool:
    """Whether the role family's resume or `.docx` cover letter holds `[TITLE]`."""
    documents = (
        _placeholder_document(source)
        for source in (family.resume, family.cover_letter)
        if source is not None
    )
    return any(
        document is not None and holds_placeholder(document, TITLE)
        for document in documents
    )


def ensure_head(
    config: Config,
    job: db.PacketJob,
    copies: list[PacketCopy],
    *,
    refresh_posting: bool = False,
) -> None:
    """Complete the folder, the job description, and the resume and cover letter copies, with their placeholders filled.

    A `.docx` cover letter copy also gets its PDF, whichever role family the copy came from. The
    copies, their PDFs, and a hand-filled job description are never replaced. The row's description replaces an existing `job_posting.md`
    only when `refresh_posting` is set.

    The copies are made only while the resume copy (made last) is missing, so a run that died
    part-way through a new head completes it, and a posting whose role family changed after its
    head was built gets nothing from the new role family (PRD 04, Edge cases).
    """
    directory = packet_paths(config, job)[0]
    directory.mkdir(parents=True, exist_ok=True)
    posting = directory / JOB_POSTING
    if job.jd_markdown is not None and (refresh_posting or not posting.exists()):
        posting.write_text(job.jd_markdown, encoding="utf-8")
    if not copies[-1].path.exists():
        for copy in copies:
            # A copy is the user's working file, so an existing one is never replaced.
            if not copy.path.exists():
                _copy_filled(copy.source, copy.path, job)
    # Not looked up in `copies`: a head built for an earlier role family may hold a cover letter this one lacks.
    for path in directory.glob(f"{glob.escape(cover_letter_stem(config, job))}.*"):
        if path.suffix.lower() == ".docx":
            ensure_pdf(path)


def build_packets(posting_id: int | None, *, dry_run: bool) -> None:
    sources = packet_sources()
    config = load_config()
    conn = db.connect()
    queue = db.packet_queue(conn, posting_id)
    if not queue:
        print("No postings awaiting a packet.")
        return
    for job in with_role_family(queue, sources.role_families):
        directory = packet_paths(config, job)[0]
        if directory.exists():
            print(f"skipped (folder exists): {directory}")
            continue
        copies = packet_copies(config, job, sources)
        print(f"{'would create' if dry_run else 'created'}: {directory}")
        if dry_run:
            for copy in copies:
                print(f"  {copy.path.name} from {copy.source}")
        else:
            ensure_head(config, job, copies)
