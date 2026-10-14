"""The redline document (PRD 04 "Redline"): Word tracked changes and comments written into a resume copy.

Pure over the loaded document (XC-9). Text offsets count the characters of a paragraph's runs, which
is also the text the redline agent is shown, so an edit's position means the same thing throughout.
"""

import copy
import itertools
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, datetime

from docx.document import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.text.paragraph import Paragraph

AUTHOR = "Claude (Redline)"
_INITIALS = "CR"
_TEXT_TAGS = ("w:br", "w:cr", "w:noBreakHyphen", "w:ptab", "w:t", "w:tab")
TEXT_XPATH = " | ".join(_TEXT_TAGS)
_TEXT_QNAMES = {qn(tag) for tag in _TEXT_TAGS}
_UNRESOLVED_XPATH = ".//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


@dataclass(frozen=True)
class TrackedChange:
    """Replace `[start, end)` of a body paragraph's text with `text`; `start == end` inserts."""

    paragraph: int
    start: int
    end: int
    text: str


@dataclass(frozen=True)
class TrackedEdit:
    """The tracked changes of one edit, which share a single comment."""

    changes: list[TrackedChange]
    comment: str


@dataclass(frozen=True)
class ParagraphComment:
    """A comment on a whole body paragraph that changes none of its text."""

    paragraph: int
    comment: str


def runs(paragraph: BaseOxmlElement) -> list[BaseOxmlElement]:
    # Runs inside a text box belong to the box's own paragraphs, not to this one (PRD 04, "Text in scope").
    return paragraph.xpath("./descendant::w:r[not(ancestor::w:txbxContent)]")


def _run_text(run: BaseOxmlElement) -> str:
    return "".join(str(child) for child in run.xpath(TEXT_XPATH))


def paragraph_text(paragraph: BaseOxmlElement) -> str:
    return "".join(_run_text(run) for run in runs(paragraph))


def body_paragraph_texts(document: Document) -> list[str]:
    """The plain text of every body paragraph, indexed as the redline's edits address them."""
    return [paragraph_text(p._p) for p in document.paragraphs]


def _wholly_bold(paragraph: Paragraph) -> bool:
    texts = [(run.bold, run.text) for run in paragraph.runs if run.text.strip()]
    return bool(texts) and all(bold is True for bold, _ in texts)


def bold_paragraphs(document: Document) -> set[int]:
    """The indices of the wholly bold body paragraphs, which the redline's rendering tags `bold`."""
    return {
        index
        for index, paragraph in enumerate(document.paragraphs)
        if _wholly_bold(paragraph)
    }


def _ends_skills_section(paragraph: Paragraph) -> bool:
    return (paragraph.style.name or "").startswith("Heading") or _wholly_bold(paragraph)


def skills_section(document: Document) -> range | None:
    """The indices of the skills section's paragraphs (PRD 04 "Skills section"); None unless one paragraph is headed `Skills`."""
    paragraphs = document.paragraphs
    headings = [
        index
        for index, paragraph in enumerate(paragraphs)
        if paragraph_text(paragraph._p).strip().rstrip(":").casefold() == "skills"
    ]
    if len(headings) != 1:
        return None
    start = headings[0] + 1
    end = next(
        (
            index
            for index in range(start, len(paragraphs))
            if _ends_skills_section(paragraphs[index])
        ),
        len(paragraphs),
    )
    return range(start, end)


def has_unresolved_changes(document: Document) -> bool:
    return bool(document.element.body.xpath(_UNRESOLVED_XPATH))


def _placed_runs(paragraph: BaseOxmlElement) -> list[tuple[BaseOxmlElement, int, int]]:
    """Each run that carries text, with its start and end offsets in the paragraph."""
    placed, offset = [], 0
    for run in runs(paragraph):
        length = len(_run_text(run))
        if length:
            placed.append((run, offset, offset + length))
        offset += length
    return placed


def content(run: BaseOxmlElement) -> list[BaseOxmlElement]:
    return [child for child in run if child.tag != qn("w:rPr")]


def text_length(element: BaseOxmlElement) -> int:
    return len(str(element)) if element.tag in _TEXT_QNAMES else 0


def set_text(element: BaseOxmlElement, text: str) -> None:
    element.text = text
    element.set(_XML_SPACE, "preserve")


def _split_run(run: BaseOxmlElement, offset: int) -> None:
    """Split a run in two at a text offset inside it; both halves keep its formatting."""
    tail = copy.deepcopy(run)
    run.addnext(tail)
    position = 0
    for head_child, tail_child in zip(content(run), content(tail), strict=True):
        length = text_length(head_child)
        end = position + length
        if length and position < offset < end:
            text = head_child.text
            set_text(head_child, text[: offset - position])
            set_text(tail_child, text[offset - position :])
        elif end <= offset and (length or position < offset):
            tail.remove(tail_child)
        else:
            run.remove(head_child)
        position = end


def _cut(paragraph: BaseOxmlElement, offset: int) -> None:
    """Make `offset` fall on a run boundary."""
    for run, start, end in _placed_runs(paragraph):
        if start < offset < end:
            _split_run(run, offset - start)
            return


class _Revisions:
    """Builds tracked-change elements with ids no other element uses, and one timestamp."""

    def __init__(self, document: Document, comment_count: int) -> None:
        used = {int(value) for value in document.element.xpath("//@w:id")}
        used.update(comment.comment_id for comment in document.comments)
        # The comments added later take the ids just above the existing ones; these start past them.
        self._ids = itertools.count(max(used, default=-1) + 1 + comment_count)
        self._now = datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")

    def element(self, tag: str) -> BaseOxmlElement:
        element = OxmlElement(tag)
        element.set(qn("w:id"), str(next(self._ids)))
        element.set(qn("w:author"), AUTHOR)
        element.set(qn("w:date"), self._now)
        return element


def _delete(
    runs: list[BaseOxmlElement], revisions: _Revisions
) -> list[BaseOxmlElement]:
    deletions: list[BaseOxmlElement] = []
    for run in runs:
        if not deletions or run.getprevious() is not deletions[-1]:
            deletion = revisions.element("w:del")
            run.addprevious(deletion)
            deletions.append(deletion)
        deletions[-1].append(run)
        for text in run.xpath("w:t"):
            text.tag = qn("w:delText")
            text.set(_XML_SPACE, "preserve")
    return deletions


def _insertion(
    text: str, format_source: BaseOxmlElement, revisions: _Revisions
) -> BaseOxmlElement:
    insertion = revisions.element("w:ins")
    run = OxmlElement("w:r")
    properties = format_source.find(qn("w:rPr"))
    if properties is not None:
        run.append(copy.deepcopy(properties))
    content = OxmlElement("w:t")
    set_text(content, text)
    run.append(content)
    insertion.append(run)
    return insertion


def _apply(
    paragraph: BaseOxmlElement, change: TrackedChange, revisions: _Revisions
) -> list[BaseOxmlElement]:
    """Write one change into the paragraph; returns its tracked-change elements in order."""
    if change.start < change.end:
        _cut(paragraph, change.start)
        _cut(paragraph, change.end)
        replaced = [
            run
            for run, start, end in _placed_runs(paragraph)
            if change.start <= start and end <= change.end
        ]
        elements = _delete(replaced, revisions)
        if change.text:
            insertion = _insertion(change.text, replaced[0], revisions)
            elements[-1].addnext(insertion)
            elements.append(insertion)
        return elements
    _cut(paragraph, change.start)
    placed = _placed_runs(paragraph)
    before = [run for run, _, end in placed if end == change.start]
    if before:
        insertion = _insertion(change.text, before[-1], revisions)
        before[-1].addnext(insertion)
    else:
        after = next(run for run, start, _ in placed if start == change.start)
        insertion = _insertion(change.text, after, revisions)
        after.addprevious(insertion)
    return [insertion]


def _anchor_comment(
    document: Document,
    first: BaseOxmlElement,
    last: BaseOxmlElement,
    text: str,
) -> None:
    comment = document.comments.add_comment(
        text=text, author=AUTHOR, initials=_INITIALS
    )

    def mark(tag: str) -> BaseOxmlElement:
        element = OxmlElement(tag)
        element.set(qn("w:id"), str(comment.comment_id))
        return element

    start, end = mark("w:commentRangeStart"), mark("w:commentRangeEnd")
    reference = OxmlElement("w:r")
    reference.append(mark("w:commentReference"))
    # The range wraps the tracked elements, so rejecting a change leaves the comment on the original text.
    first.addprevious(start)
    last.addnext(end)
    end.addnext(reference)


def _comment_paragraph(
    document: Document, paragraph: BaseOxmlElement, text: str
) -> None:
    """Anchor a comment on everything in the paragraph, which is left as it is."""
    if not paragraph.xpath("./*[not(self::w:pPr)]"):
        # A range needs something to wrap; an empty run changes no text.
        paragraph.append(OxmlElement("w:r"))
    children = paragraph.xpath("./*[not(self::w:pPr)]")
    _anchor_comment(document, children[0], children[-1], text)


def apply_edits(
    document: Document,
    edits: list[TrackedEdit],
    comments: Sequence[ParagraphComment] = (),
) -> None:
    """Write the edits into the document as tracked changes, one comment around each edit.

    Later text is changed first, so the offsets of the changes still to apply stay valid. Each
    paragraph comment is anchored afterwards, around the paragraph's final content.
    """
    revisions = _Revisions(document, len(edits) + len(comments))
    paragraphs = document.paragraphs
    for edit in sorted(
        edits, key=lambda e: (e.changes[0].paragraph, e.changes[0].start), reverse=True
    ):
        # Changes apply last to first, so the earliest change's elements come last.
        applied = [
            _apply(paragraphs[change.paragraph]._p, change, revisions)
            for change in sorted(
                edit.changes, key=lambda c: (c.start, c.end), reverse=True
            )
        ]
        _anchor_comment(document, applied[-1][0], applied[0][-1], edit.comment)
    for comment in comments:
        _comment_paragraph(document, paragraphs[comment.paragraph]._p, comment.comment)
