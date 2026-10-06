"""`jsa refetch` (issue #14; PRD 04 "Reconciliation" and "Projection reads and writes"; PRD 02 "Refetch scope" and "JD capture"; XC-4, XC-5, XC-6).

The database is the libSQL container (and a `file:` database). The outside world is replaced at the
documented seams: HTTP at the transport (the employer's ATS record), Claude at `agent_loop.query`, and
`gws` by a stub script reached through `JSA_GWS_BIN`. The profile and the packets directory are temporary
directories.
"""

import html
import itertools
import json
import sys
from pathlib import Path

import docx
import httpx
import pytest
from conftest import drop_all_tables
from pandoc_helpers import install_pandoc
from profile_helpers import copy_example, write_config_toml
from test_generate import (
    BASE_LINE,
    CHECKLIST_TEXT,
    CONFIG,
    LONG_AGO,
    PDF,
    Agent,
    column,
    edit_copy,
    failed,
    jsa,
)
from test_review import Web
from test_tracker import params_of

from jsa import agent_loop, db, tracker
from jsa.naming import title_slug

STUB = """\
#!{python}
import json, os, sys

argv = sys.argv[1:]
with open(os.environ["STUB_LOG"], "a") as log:
    log.write(json.dumps(argv) + "\\n")
method = next(word for word in argv if word in ("get", "update", "append"))
if method == "get":
    if os.environ.get("STUB_GET_FAIL"):
        print("boom", file=sys.stderr)
        sys.exit(1)
    with open(os.environ["STUB_SHEET"]) as sheet:
        print(json.dumps({{"values": json.load(sheet)}}))
elif method == "update":
    if os.environ.get("STUB_UPDATE_FAIL"):
        print("boom", file=sys.stderr)
        sys.exit(1)
    print(json.dumps({{"updatedCells": 1}}))
else:
    print(json.dumps({{"updates": {{"updatedRows": 1}}}}))
"""

OLD_TITLE = "Staff Engineer"
NEW_TITLE = "Principal Engineer"
OLD_JD = "Old description of the role."
NEW_JD = "Brand new description of the Frobnicator role."
OLD_LOCATION = "Boston, MA"
NEW_LOCATION = "Remote, US"
OLD_CHECKLIST = "OLD CHECKLIST"
REVISION = "USER REVISION LINE"
OLD_DIR = "Acme Widgets - Staff Engineer"
NEW_DIR = "Acme Widgets - Principal Engineer"
OLD_COPY = "PatExample_Resume_StaffEngineer_AcmeWidgets.docx"
NEW_COPY = "PatExample_Resume_PrincipalEngineer_AcmeWidgets.docx"
APPLIED = "2026-10-02"
JOB_IDS = itertools.count(4000000)


@pytest.fixture(autouse=True)
def agent(monkeypatch):
    stand_in = Agent()
    monkeypatch.setattr(agent_loop, "query", stand_in.query)
    return stand_in


@pytest.fixture(autouse=True)
def web(monkeypatch):
    """Pages the test doesn't route answer 200 with plain HTML; nothing reaches the network."""
    stand_in = Web()

    def handle(_transport, request):
        return stand_in.handle(request)

    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", handle)
    return stand_in


@pytest.fixture(autouse=True)
def env(tmp_path, monkeypatch):
    """A profile with a distinctive base resume, and its packets directory."""
    profile = copy_example(tmp_path / "profile")
    packets = tmp_path / "packets"
    write_config_toml(profile, CONFIG.format(packets=packets))
    document = docx.Document()
    document.add_paragraph().add_run("Riley Resumeperson").bold = True
    document.add_paragraph(BASE_LINE)
    document.save(profile / "resume.docx")
    monkeypatch.setenv("JSA_PROFILE_DIR", str(profile))
    monkeypatch.delenv("JSA_GENERATE_WORKERS", raising=False)
    return profile, packets


@pytest.fixture(autouse=True)
def gws(tmp_path, monkeypatch):
    """A stub `gws` that logs every call and serves the Sheet written by `set_sheet`."""
    stub = tmp_path / "stub-gws"
    stub.write_text(STUB.format(python=sys.executable), encoding="utf-8")
    stub.chmod(0o755)
    log = tmp_path / "gws.log"
    sheet = tmp_path / "sheet.json"
    sheet.write_text("[]", encoding="utf-8")
    monkeypatch.setenv("JSA_GWS_BIN", str(stub))
    monkeypatch.setenv("STUB_LOG", str(log))
    monkeypatch.setenv("STUB_SHEET", str(sheet))
    for name in ("STUB_GET_FAIL", "STUB_UPDATE_FAIL"):
        monkeypatch.delenv(name, raising=False)
    return log


@pytest.fixture(autouse=True)
def pandoc(tmp_path, monkeypatch):
    """A stub `pandoc` whose invocations are logged; returns the log's path."""
    return install_pandoc(tmp_path, monkeypatch)


@pytest.fixture
def rdb(db_url):
    """A connection to a database holding no postings but the test's own."""
    drop_all_tables(db_url)
    connection = db.connect()
    yield connection
    connection.close()


def set_sheet(gws, *rows):
    """The Sheet as `gws` returns it: a header, then one row per `(posting id, Date Applied)`."""
    cells = [["ID", "Company", "Title", "URL", "Posted", "Added", "Applied", "Status"]]
    for posting_id, applied in rows:
        cells.append(
            [str(posting_id), "Acme Widgets", OLD_TITLE, "u", "", "2026-10-01"]
            + ([applied] if applied else [])
        )
    gws.with_name("sheet.json").write_text(json.dumps(cells), encoding="utf-8")


def seed(
    conn,
    *,
    title=OLD_TITLE,
    decision="Apply",
    search_agent="claude",
    jd=OLD_JD,
    location=OLD_LOCATION,
    company="Acme Widgets, Inc.",
):
    job_id = next(JOB_IDS)
    posting_id = db.insert_posting(
        conn,
        company=company,
        title=title,
        url=f"https://job-boards.greenhouse.io/acme-widgets/jobs/{job_id}",
        search_agent=search_agent,
    )
    conn.execute(
        "UPDATE postings SET decision = ?, decided_at = ?, jd_markdown = ?, location = ? "
        "WHERE id = ?",
        (decision, LONG_AGO if decision else None, jd, location, posting_id),
    )
    return posting_id


def _record_url(conn, posting_id) -> str:
    job_id = column(conn, posting_id, "url").rsplit("/", 1)[1]
    return f"https://boards-api.greenhouse.io/v1/boards/acme-widgets/jobs/{job_id}"


def employer(
    web, conn, posting_id, *, title=NEW_TITLE, body=NEW_JD, location=NEW_LOCATION
):
    """The employer's ATS record for the posting, as it reads now."""
    web.routes[_record_url(conn, posting_id)] = httpx.Response(
        200,
        json={
            "title": title,
            "content": html.escape(f"<p>{body}</p>"),
            "location": {"name": location},
        },
    )


def employer_down(web, conn, posting_id):
    web.routes[_record_url(conn, posting_id)] = httpx.Response(500)


def whole_row(conn, posting_id):
    return conn.execute("SELECT * FROM postings WHERE id = ?", (posting_id,)).fetchone()


def make_packet(conn, packets, posting_id, *, directory=OLD_DIR, copy=OLD_COPY):
    """A packet as `jsa generate` leaves it, with the user's revision in the resume copy and a note of their own."""
    folder = packets / directory
    folder.mkdir(parents=True)
    (folder / "job_posting.md").write_text(
        column(conn, posting_id, "jd_markdown"), encoding="utf-8"
    )
    (folder / "resume_checklist.md").write_text(OLD_CHECKLIST, encoding="utf-8")
    (folder / "notes.txt").write_text("my own notes", encoding="utf-8")
    document = docx.Document()
    document.add_paragraph(BASE_LINE)
    document.save(folder / copy)
    edit_copy(folder / copy, REVISION)
    return folder


def tree(root: Path) -> dict[str, bytes]:
    """Every file under `root`, by relative path."""
    return {
        str(path.relative_to(root)): path.read_bytes()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def stamps(root: Path) -> dict[str, int]:
    return {
        str(path.relative_to(root)): path.stat().st_mtime_ns
        for path in sorted(root.rglob("*"))
    }


def refetch(monkeypatch, capsys, *args):
    return jsa(monkeypatch, capsys, "refetch", *args)


def title_updates(log: Path) -> list[tuple[str, str]]:
    """The `(range, value)` of every Sheet cell write the stub saw."""
    writes = []
    if not log.exists():
        return writes
    for line in log.read_text().splitlines():
        argv = json.loads(line)
        if "update" in argv:
            params = params_of(argv)
            writes.append((params["range"], params["values"][0][0]))
    return writes


def flagged(code, output, posting_id) -> bool:
    return code != 0 or str(posting_id) in output or "flag" in output.lower()


# --- the tracker index --------------------------------------------------------


def test_the_index_maps_posting_ids_to_row_numbers_and_date_applied(gws):
    set_sheet(gws, (7, ""), (12, APPLIED))
    index = tracker.sheet_index("sheet-123")
    assert index[7].number == 2 and index[7].date_applied == ""
    assert index[12].number == 3 and index[12].date_applied == APPLIED


def test_the_index_skips_rows_whose_id_cell_is_not_an_integer(gws):
    cells = [
        ["ID", "Company"],
        ["abc", "Acme"],
        [],
        ["", "Acme"],
        ["4.5", "Acme"],
        ["21", "Acme", "Role", "u", "", "2026-10-01", APPLIED],
    ]
    gws.with_name("sheet.json").write_text(json.dumps(cells), encoding="utf-8")
    index = tracker.sheet_index("sheet-123")
    assert list(index) == [21]
    assert index[21].number == 6


def test_a_title_is_written_to_the_title_cell_of_the_rows_number(gws):
    tracker.set_title("sheet-123", 9, "Principal Engineer")
    assert title_updates(gws) == [(f"{tracker.TAB}!C9", "Principal Engineer")]


@pytest.mark.parametrize("lead", ["=", "+", "-", "@"])
def test_a_title_that_looks_like_a_formula_is_written_as_literal_text(gws, lead):
    text = f"{lead}HYPERLINK(1,2)"
    tracker.set_title("sheet-123", 3, text)
    [(_, value)] = title_updates(gws)
    assert value != text and value.startswith("'") and text in value


def test_a_failed_title_write_raises(gws, monkeypatch):
    monkeypatch.setenv("STUB_UPDATE_FAIL", "1")
    with pytest.raises(Exception, match="gws"):
        tracker.set_title("sheet-123", 3, "Principal Engineer")


# --- scope --------------------------------------------------------------------


def test_the_default_scope_is_unapplied_apply_rows(rdb, web, gws, monkeypatch, capsys):
    absent = seed(rdb)
    blank = seed(rdb)
    applied = seed(rdb)
    skipped = seed(rdb, decision="Skip")
    undecided = seed(rdb, decision=None)
    for posting_id in (absent, blank, applied, skipped, undecided):
        employer(web, rdb, posting_id)
    set_sheet(gws, (blank, ""), (applied, APPLIED), (skipped, ""))
    code, _ = refetch(monkeypatch, capsys)
    assert code == 0
    for posting_id in (absent, blank):
        assert column(rdb, posting_id, "jd_markdown") == NEW_JD
    for posting_id in (applied, skipped, undecided):
        assert column(rdb, posting_id, "jd_markdown") == OLD_JD
        assert column(rdb, posting_id, "title") == OLD_TITLE


def test_a_failed_sheet_read_in_the_default_scope_exits_non_zero_before_any_change(
    rdb, web, gws, monkeypatch, capsys
):
    ids = [seed(rdb), seed(rdb)]
    for posting_id in ids:
        employer(web, rdb, posting_id)
    before = [whole_row(rdb, posting_id) for posting_id in ids]
    monkeypatch.setenv("STUB_GET_FAIL", "1")
    code, _ = refetch(monkeypatch, capsys)
    assert code != 0
    assert [whole_row(rdb, posting_id) for posting_id in ids] == before
    assert title_updates(gws) == []


def test_all_processes_every_row_including_applied_and_skipped(
    rdb, web, gws, monkeypatch, capsys
):
    applied = seed(rdb)
    skipped = seed(rdb, decision="Skip")
    undecided = seed(rdb, decision=None)
    for posting_id in (applied, skipped, undecided):
        employer(web, rdb, posting_id)
    set_sheet(gws, (applied, APPLIED))
    code, _ = refetch(monkeypatch, capsys, "--all")
    assert code == 0
    for posting_id in (applied, skipped, undecided):
        assert column(rdb, posting_id, "jd_markdown") == NEW_JD
    assert title_updates(gws) == []


def test_all_continues_when_the_sheet_read_fails(rdb, web, gws, monkeypatch, capsys):
    posting_id = seed(rdb)
    employer(web, rdb, posting_id)
    monkeypatch.setenv("STUB_GET_FAIL", "1")
    refetch(monkeypatch, capsys, "--all")
    assert column(rdb, posting_id, "jd_markdown") == NEW_JD
    assert column(rdb, posting_id, "title") == NEW_TITLE
    assert title_updates(gws) == []


def test_id_processes_that_one_row_unconditionally(rdb, web, gws, monkeypatch, capsys):
    target = seed(rdb)
    other = seed(rdb)
    skipped = seed(rdb, decision="Skip")
    for posting_id in (target, other, skipped):
        employer(web, rdb, posting_id)
    set_sheet(gws, (target, APPLIED), (other, ""))
    code, _ = refetch(monkeypatch, capsys, "--id", target)
    assert code == 0
    assert column(rdb, target, "jd_markdown") == NEW_JD
    assert column(rdb, other, "jd_markdown") == OLD_JD
    code, _ = refetch(monkeypatch, capsys, "--id", skipped)
    assert code == 0
    assert column(rdb, skipped, "jd_markdown") == NEW_JD


def test_id_continues_when_the_sheet_read_fails(rdb, web, gws, monkeypatch, capsys):
    posting_id = seed(rdb)
    employer(web, rdb, posting_id)
    monkeypatch.setenv("STUB_GET_FAIL", "1")
    refetch(monkeypatch, capsys, "--id", posting_id)
    assert column(rdb, posting_id, "jd_markdown") == NEW_JD
    assert title_updates(gws) == []


# --- the insert rule ----------------------------------------------------------


def test_a_searched_row_takes_the_employers_title_slug_description_and_location(
    rdb, web, monkeypatch, capsys
):
    posting_id = seed(rdb)
    employer(web, rdb, posting_id)
    code, _ = refetch(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert column(rdb, posting_id, "title") == NEW_TITLE
    assert column(rdb, posting_id, "title_slug") == title_slug(NEW_TITLE)
    assert NEW_JD in column(rdb, posting_id, "jd_markdown")
    assert column(rdb, posting_id, "location") == NEW_LOCATION


def test_a_manual_row_keeps_its_title_and_slug_but_takes_description_and_location(
    rdb, web, monkeypatch, capsys
):
    posting_id = seed(rdb, search_agent="manual")
    slug = column(rdb, posting_id, "title_slug")
    employer(web, rdb, posting_id)
    code, _ = refetch(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert column(rdb, posting_id, "title") == OLD_TITLE
    assert column(rdb, posting_id, "title_slug") == slug
    assert NEW_JD in column(rdb, posting_id, "jd_markdown")
    assert column(rdb, posting_id, "location") == NEW_LOCATION


def test_a_manual_row_is_never_retitled_in_the_sheet_or_renamed_on_disk(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb, search_agent="manual")
    make_packet(rdb, packets, posting_id)
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    refetch(monkeypatch, capsys, "--id", posting_id)
    assert title_updates(gws) == []
    assert {path.name for path in packets.iterdir()} == {OLD_DIR}


@pytest.mark.parametrize("fault", ["server-error", "unparseable"])
def test_a_failed_fetch_leaves_every_column_and_file_and_cell_untouched(
    rdb, web, gws, env, agent, fault, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    make_packet(rdb, packets, posting_id)
    set_sheet(gws, (posting_id, ""))
    if fault == "server-error":
        employer_down(web, rdb, posting_id)
    else:
        web.routes[_record_url(rdb, posting_id)] = httpx.Response(
            200, text="not json at all"
        )
    row, files = whole_row(rdb, posting_id), tree(packets)
    refetch(monkeypatch, capsys, "--id", posting_id)
    assert whole_row(rdb, posting_id) == row
    assert tree(packets) == files
    assert title_updates(gws) == []
    assert agent.calls == []


def test_one_failed_fetch_does_not_stop_the_other_rows(
    rdb, web, gws, monkeypatch, capsys
):
    broken, fine = seed(rdb), seed(rdb)
    employer_down(web, rdb, broken)
    employer(web, rdb, fine)
    set_sheet(gws)
    refetch(monkeypatch, capsys)
    assert column(rdb, broken, "jd_markdown") == OLD_JD
    assert column(rdb, fine, "jd_markdown") == NEW_JD


# --- title propagation --------------------------------------------------------


def test_a_retitled_unapplied_tracked_row_gets_its_title_cell_rewritten(
    rdb, web, gws, monkeypatch, capsys
):
    seed(rdb)
    posting_id = seed(rdb)
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id - 1, APPLIED), (posting_id, ""))
    code, _ = refetch(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert title_updates(gws) == [(f"{tracker.TAB}!C3", NEW_TITLE)]


def test_a_retitle_to_formula_text_is_written_to_the_sheet_as_literal_text(
    rdb, web, gws, monkeypatch, capsys
):
    posting_id = seed(rdb)
    employer(web, rdb, posting_id, title='=HYPERLINK("http://x","y")')
    set_sheet(gws, (posting_id, ""))
    refetch(monkeypatch, capsys, "--id", posting_id)
    [(cell, value)] = title_updates(gws)
    assert cell.endswith("C2") and value.startswith("'=HYPERLINK")


def test_a_retitled_row_with_a_date_applied_keeps_its_title_cell(
    rdb, web, gws, monkeypatch, capsys
):
    posting_id = seed(rdb)
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, APPLIED))
    refetch(monkeypatch, capsys, "--all")
    assert column(rdb, posting_id, "title") == NEW_TITLE
    assert title_updates(gws) == []


def test_a_retitled_row_missing_from_the_sheet_has_no_cell_to_rewrite(
    rdb, web, gws, monkeypatch, capsys
):
    posting_id = seed(rdb)
    employer(web, rdb, posting_id)
    set_sheet(gws)
    code, _ = refetch(monkeypatch, capsys)
    assert code == 0
    assert column(rdb, posting_id, "title") == NEW_TITLE
    assert title_updates(gws) == []


def test_an_unchanged_title_is_not_rewritten_in_the_sheet(
    rdb, web, gws, monkeypatch, capsys
):
    posting_id = seed(rdb)
    employer(web, rdb, posting_id, title=OLD_TITLE)
    set_sheet(gws, (posting_id, ""))
    refetch(monkeypatch, capsys, "--id", posting_id)
    assert title_updates(gws) == []
    assert NEW_JD in column(rdb, posting_id, "jd_markdown")


def test_a_failed_sheet_write_is_flagged_and_the_database_update_stands(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    make_packet(rdb, packets, posting_id)
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    monkeypatch.setenv("STUB_UPDATE_FAIL", "1")
    code, output = refetch(monkeypatch, capsys, "--id", posting_id)
    assert flagged(code, output, posting_id)
    assert column(rdb, posting_id, "title") == NEW_TITLE
    assert NEW_JD in column(rdb, posting_id, "jd_markdown")
    # The failed write doesn't block the packet's refresh either.
    assert (packets / NEW_DIR / NEW_COPY).exists()


# --- the packet ---------------------------------------------------------------


def test_a_retitled_row_with_a_packet_is_renamed_and_regenerated_in_place(
    rdb, web, gws, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    folder = make_packet(rdb, packets, posting_id)
    revised = (folder / OLD_COPY).read_bytes()
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    code, _ = refetch(monkeypatch, capsys)
    assert code == 0
    assert {path.name for path in packets.iterdir()} == {NEW_DIR}
    renamed = packets / NEW_DIR
    assert (renamed / NEW_COPY).read_bytes() == revised
    assert not (renamed / OLD_COPY).exists()
    assert NEW_JD in (renamed / "job_posting.md").read_text(encoding="utf-8")
    checklist = (renamed / "resume_checklist.md").read_text(encoding="utf-8")
    assert checklist.strip() == CHECKLIST_TEXT.strip()
    assert (renamed / "notes.txt").read_text(encoding="utf-8") == "my own notes"
    assert {path.name for path in renamed.iterdir()} == {
        "job_posting.md",
        "resume_checklist.md",
        PDF,
        "notes.txt",
        NEW_COPY,
    }


def test_the_refreshed_checklist_assesses_the_users_revised_copy_and_the_new_posting(
    rdb, web, gws, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    make_packet(rdb, packets, posting_id)
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    refetch(monkeypatch, capsys)
    [prompt] = agent.prompts
    assert REVISION in prompt
    assert NEW_TITLE in prompt and NEW_JD in prompt
    assert OLD_JD not in prompt


def test_a_description_only_change_regenerates_without_renaming(
    rdb, web, gws, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    folder = make_packet(rdb, packets, posting_id)
    revised = (folder / OLD_COPY).read_bytes()
    employer(web, rdb, posting_id, title=OLD_TITLE, location=OLD_LOCATION)
    set_sheet(gws, (posting_id, ""))
    code, _ = refetch(monkeypatch, capsys)
    assert code == 0
    assert {path.name for path in packets.iterdir()} == {OLD_DIR}
    assert (folder / OLD_COPY).read_bytes() == revised
    assert NEW_JD in (folder / "job_posting.md").read_text(encoding="utf-8")
    checklist = (folder / "resume_checklist.md").read_text(encoding="utf-8")
    assert checklist.strip() == CHECKLIST_TEXT.strip()
    assert len(agent.calls) == 1
    assert (folder / PDF).read_text(encoding="utf-8") == f"PDF OF: {checklist}"
    assert title_updates(gws) == []


def test_a_location_only_change_touches_no_file(
    rdb, web, gws, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    employer(web, rdb, posting_id, title=OLD_TITLE, location=OLD_LOCATION)
    set_sheet(gws, (posting_id, ""))
    refetch(monkeypatch, capsys, "--id", posting_id)
    make_packet(rdb, packets, posting_id)
    agent.calls.clear()
    employer(web, rdb, posting_id, title=OLD_TITLE, location=NEW_LOCATION)
    files, times = tree(packets), stamps(packets)
    code, _ = refetch(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert column(rdb, posting_id, "location") == NEW_LOCATION
    assert tree(packets) == files and stamps(packets) == times
    assert agent.calls == []
    assert title_updates(gws) == []


def test_an_unchanged_posting_touches_nothing_and_a_second_run_is_a_no_op(
    rdb, web, gws, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    refetch(monkeypatch, capsys, "--id", posting_id)
    make_packet(rdb, packets, posting_id, directory=NEW_DIR, copy=NEW_COPY)
    agent.calls.clear()
    row, files, times = whole_row(rdb, posting_id), tree(packets), stamps(packets)
    writes = len(title_updates(gws))
    code, _ = refetch(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert whole_row(rdb, posting_id) == row
    assert tree(packets) == files and stamps(packets) == times
    assert agent.calls == []
    assert len(title_updates(gws)) == writes


@pytest.mark.parametrize("change", ["description", "title", "location"])
def test_a_row_without_a_packet_directory_gets_none(
    rdb, web, gws, env, agent, change, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    employer(
        web,
        rdb,
        posting_id,
        title=NEW_TITLE if change == "title" else OLD_TITLE,
        body=NEW_JD if change == "description" else OLD_JD,
        location=NEW_LOCATION if change == "location" else OLD_LOCATION,
    )
    set_sheet(gws, (posting_id, ""))
    code, _ = refetch(monkeypatch, capsys)
    assert code == 0
    assert not packets.exists() or list(packets.iterdir()) == []
    assert agent.calls == []


def test_a_taken_directory_name_blocks_the_rename_and_flags_the_row(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    folder = make_packet(rdb, packets, posting_id)
    squatter = packets / NEW_DIR
    squatter.mkdir()
    (squatter / "theirs.txt").write_text("someone else's packet", encoding="utf-8")
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    before = {path: tree(path) for path in (folder, squatter)}
    code, output = refetch(monkeypatch, capsys)
    assert flagged(code, output, posting_id)
    assert {path.name for path in packets.iterdir()} == {OLD_DIR, NEW_DIR}
    assert tree(squatter) == before[squatter]
    assert (folder / OLD_COPY).read_bytes() == before[folder][OLD_COPY]
    for name in before[folder]:
        assert (folder / name).exists()
    assert (folder / "notes.txt").read_bytes() == before[folder]["notes.txt"]


def test_a_taken_resume_file_name_blocks_the_rename_and_flags_the_row(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    folder = make_packet(rdb, packets, posting_id)
    squatter = b"a different file at the new resume name"
    (folder / NEW_COPY).write_bytes(squatter)
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    original = (folder / OLD_COPY).read_bytes()
    code, output = refetch(monkeypatch, capsys)
    assert flagged(code, output, posting_id)
    assert {path.name for path in packets.iterdir()} == {OLD_DIR}
    assert (folder / OLD_COPY).read_bytes() == original
    assert (folder / NEW_COPY).read_bytes() == squatter
    assert (folder / "notes.txt").exists()


def test_a_retitle_that_frees_a_shared_name_renames_the_other_postings_packet(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    first, second = seed(rdb), seed(rdb)
    make_packet(rdb, packets, first)
    second_folder = make_packet(rdb, packets, second, directory=f"{OLD_DIR} ({second})")
    second_files = tree(second_folder)
    employer(web, rdb, first)
    set_sheet(gws, (first, ""))
    code, _ = refetch(monkeypatch, capsys, "--id", first)
    assert code == 0
    assert {path.name for path in packets.iterdir()} == {NEW_DIR, OLD_DIR}
    assert (packets / NEW_DIR / NEW_COPY).exists()
    assert tree(packets / OLD_DIR) == second_files


def test_a_retitle_onto_another_postings_name_gives_that_posting_the_suffix(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    first = seed(rdb, title=NEW_TITLE)
    second = seed(rdb)
    make_packet(rdb, packets, first, directory=NEW_DIR, copy=NEW_COPY)
    second_folder = make_packet(rdb, packets, second)
    second_files = tree(second_folder)
    employer(web, rdb, first, title=OLD_TITLE)
    set_sheet(gws, (first, ""))
    refetch(monkeypatch, capsys, "--id", first)
    suffixed = packets / f"{OLD_DIR} ({second})"
    assert tree(suffixed) == second_files


def test_no_packet_file_is_lost_when_other_postings_packets_are_reconciled(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    first, second = seed(rdb), seed(rdb)
    make_packet(rdb, packets, first)
    make_packet(rdb, packets, second, directory=f"{OLD_DIR} ({second})")
    notes = sorted(path.read_text() for path in packets.rglob("notes.txt"))
    revised = {path.read_bytes() for path in packets.rglob("*.docx")}
    employer(web, rdb, first)
    set_sheet(gws, (first, ""))
    refetch(monkeypatch, capsys, "--id", first)
    assert sorted(path.read_text() for path in packets.rglob("notes.txt")) == notes
    assert revised <= {path.read_bytes() for path in packets.rglob("*.docx")}


def test_the_other_postings_resume_copy_keeps_its_name_and_the_user_revision(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    first, second = seed(rdb), seed(rdb)
    make_packet(rdb, packets, first)
    second_folder = make_packet(rdb, packets, second, directory=f"{OLD_DIR} ({second})")
    revised = (second_folder / OLD_COPY).read_bytes()
    employer(web, rdb, first)
    set_sheet(gws, (first, ""))
    refetch(monkeypatch, capsys, "--id", first)
    assert (packets / OLD_DIR / OLD_COPY).read_bytes() == revised


def test_a_taken_name_leaves_the_other_postings_packet_in_place_and_flags_the_run(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    first, second = seed(rdb), seed(rdb)
    first_folder = make_packet(rdb, packets, first)
    suffixed = f"{OLD_DIR} ({second})"
    second_folder = make_packet(rdb, packets, second, directory=suffixed)
    squatter = packets / NEW_DIR
    squatter.mkdir()
    (squatter / "theirs.txt").write_text("someone else's packet", encoding="utf-8")
    before = {path: tree(path) for path in (first_folder, second_folder, squatter)}
    employer(web, rdb, first)
    set_sheet(gws, (first, ""))
    code, _ = refetch(monkeypatch, capsys, "--id", first)
    assert code != 0
    assert {path.name for path in packets.iterdir()} == {OLD_DIR, suffixed, NEW_DIR}
    for path, files in before.items():
        assert tree(path) == files


def test_a_posting_without_a_packet_gets_none_when_another_is_retitled(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    first = seed(rdb)
    seed(rdb)
    make_packet(rdb, packets, first)
    employer(web, rdb, first)
    set_sheet(gws, (first, ""))
    code, _ = refetch(monkeypatch, capsys, "--id", first)
    assert code == 0
    assert {path.name for path in packets.iterdir()} == {NEW_DIR}


def test_a_retitle_leaves_packets_of_other_names_and_companies_alone(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    first = seed(rdb)
    other_title = seed(rdb, title="Data Engineer")
    other_company = seed(rdb, company="Globex Corporation")
    make_packet(rdb, packets, first)
    make_packet(
        rdb,
        packets,
        other_title,
        directory="Acme Widgets - Data Engineer",
        copy="x.docx",
    )
    make_packet(
        rdb, packets, other_company, directory="Globex - Staff Engineer", copy="y.docx"
    )
    untouched = {
        name: tree(packets / name)
        for name in ("Acme Widgets - Data Engineer", "Globex - Staff Engineer")
    }
    employer(web, rdb, first)
    set_sheet(gws, (first, ""))
    code, _ = refetch(monkeypatch, capsys, "--id", first)
    assert code == 0
    for name, files in untouched.items():
        assert tree(packets / name) == files
    assert {path.name for path in packets.iterdir()} == {NEW_DIR, *untouched}


def test_a_dry_run_renames_no_other_postings_packet(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    first, second = seed(rdb), seed(rdb)
    make_packet(rdb, packets, first)
    make_packet(rdb, packets, second, directory=f"{OLD_DIR} ({second})")
    employer(web, rdb, first)
    set_sheet(gws, (first, ""))
    files, times = tree(packets), stamps(packets)
    code, _ = refetch(monkeypatch, capsys, "--dry-run", "--id", first)
    assert code == 0
    assert tree(packets) == files and stamps(packets) == times


def test_a_second_refetch_after_reconciling_other_packets_changes_nothing(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    first, second = seed(rdb), seed(rdb)
    make_packet(rdb, packets, first)
    make_packet(rdb, packets, second, directory=f"{OLD_DIR} ({second})")
    employer(web, rdb, first)
    set_sheet(gws, (first, ""))
    refetch(monkeypatch, capsys, "--id", first)
    files = tree(packets)
    code, _ = refetch(monkeypatch, capsys, "--id", first)
    assert code == 0
    assert tree(packets) == files


def test_a_failed_regenerate_is_flagged_and_nothing_is_rolled_back(
    rdb, web, gws, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    folder = make_packet(rdb, packets, posting_id)
    revised = (folder / OLD_COPY).read_bytes()
    agent.respond = lambda prompt: failed(result="overloaded")
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    code, output = refetch(monkeypatch, capsys)
    assert flagged(code, output, posting_id)
    assert column(rdb, posting_id, "title") == NEW_TITLE
    assert {path.name for path in packets.iterdir()} == {NEW_DIR}
    assert (packets / NEW_DIR / NEW_COPY).read_bytes() == revised
    assert (packets / NEW_DIR / "resume_checklist.md").exists()
    assert (packets / NEW_DIR / "notes.txt").exists()


def test_a_pandoc_failure_on_regenerate_is_flagged_and_the_new_checklist_stays(
    rdb, web, gws, env, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    make_packet(rdb, packets, posting_id)
    monkeypatch.setenv("PANDOC_FAIL", "1")
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    code, output = refetch(monkeypatch, capsys)
    assert flagged(code, output, posting_id)
    renamed = packets / NEW_DIR
    checklist = (renamed / "resume_checklist.md").read_text(encoding="utf-8")
    assert checklist.strip() == CHECKLIST_TEXT.strip()
    assert not (renamed / PDF).exists()


# --- dry run ------------------------------------------------------------------


@pytest.mark.parametrize("flags", [(), ("--all",)], ids=["default", "all"])
def test_a_dry_run_changes_no_row_file_or_cell_and_calls_no_model(
    rdb, web, gws, env, agent, flags, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    other = seed(rdb, search_agent="manual")
    make_packet(rdb, packets, posting_id)
    for target in (posting_id, other):
        employer(web, rdb, target)
    set_sheet(gws, (posting_id, ""), (other, ""))
    rows = [whole_row(rdb, target) for target in (posting_id, other)]
    files, times = tree(packets), stamps(packets)
    code, output = refetch(monkeypatch, capsys, "--dry-run", *flags)
    assert code == 0
    assert output.strip()
    assert [whole_row(rdb, target) for target in (posting_id, other)] == rows
    assert tree(packets) == files and stamps(packets) == times
    assert title_updates(gws) == []
    assert agent.calls == []


def test_a_dry_run_with_an_id_changes_nothing(
    rdb, web, gws, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(rdb)
    make_packet(rdb, packets, posting_id)
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, ""))
    row, files = whole_row(rdb, posting_id), tree(packets)
    code, _ = refetch(monkeypatch, capsys, "--dry-run", "--id", posting_id)
    assert code == 0
    assert whole_row(rdb, posting_id) == row
    assert tree(packets) == files
    assert title_updates(gws) == [] and agent.calls == []


# --- the command line and the database ----------------------------------------


@pytest.mark.parametrize("value", ["0", "-3", "abc"])
def test_a_non_positive_or_non_numeric_id_is_rejected(rdb, value, monkeypatch, capsys):
    code, _ = refetch(monkeypatch, capsys, "--id", value)
    assert code != 0


def test_refetch_never_changes_decision_or_tracking_state(
    rdb, web, gws, monkeypatch, capsys
):
    posting_id = seed(rdb)
    rdb.execute(
        "UPDATE postings SET fit_feedback = 'strong', added_to_tracker = 1 WHERE id = ?",
        (posting_id,),
    )
    employer(web, rdb, posting_id)
    set_sheet(gws, (posting_id, APPLIED))
    refetch(monkeypatch, capsys, "--all")
    assert column(rdb, posting_id, "decision") == "Apply"
    assert column(rdb, posting_id, "decided_at") == LONG_AGO
    assert column(rdb, posting_id, "fit_feedback") == "strong"
    assert column(rdb, posting_id, "added_to_tracker") == 1
    assert column(rdb, posting_id, "closed_at") is None
