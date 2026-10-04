"""`jsa packet` (issue #11; PRD 04 "Base resume" and "Packet directory"; PRD 02 "Packet queue"; XC-9, XC-11).

The database is the libSQL container (and a `file:` database); the packets directory and the profile are
temporary directories, so nothing outside the test's own paths is touched.
"""

import socket
import sqlite3
import sys
from pathlib import Path

import docx
import pytest
from conftest import drop_all_tables, unique_url
from profile_helpers import EXAMPLE_DIR, copy_example, write_config_toml

from jsa import cli, db
from jsa.naming import normalize_company, packet_dir_name, resume_file_stem, title_slug

LONG_AGO = "2020-01-01T00:00:00.000Z"
JD = "# Staff Engineer\n\nBuild the platform.\n"
PLAIN = "Acme Widgets - Staff Engineer"
RESUME_COPY = "PatExample_Resume_StaffEngineer_AcmeWidgets.docx"


@pytest.fixture
def pdb(db_url):
    """A connection to a database holding no postings but the test's own."""
    drop_all_tables(db_url)
    connection = db.connect()
    yield connection
    connection.close()


@pytest.fixture
def env(tmp_path, monkeypatch):
    """A profile with a distinctive base resume and `candidate_name`, and its packets directory."""
    profile = copy_example(tmp_path / "profile")
    packets = tmp_path / "packets"
    write_config_toml(
        profile, f'candidate_name = "Pat Example"\npackets_dir = "{packets}"\n'
    )
    document = docx.Document()
    document.add_paragraph("Pat Example, a fictional candidate")
    document.save(profile / "resume.docx")
    monkeypatch.setenv("JSA_PROFILE_DIR", str(profile))
    return profile, packets


def seed(
    conn,
    *,
    company="Acme Widgets, Inc.",
    title="Staff Engineer",
    decision="Apply",
    tracked=False,
    closed=False,
    jd=JD,
):
    posting_id = db.insert_posting(
        conn, company=company, title=title, url=unique_url(), search_agent="claude"
    )
    conn.execute(
        "UPDATE postings SET decision = ?, decided_at = ?, added_to_tracker = ?, "
        "closed_at = ?, jd_markdown = ? WHERE id = ?",
        (
            decision,
            LONG_AGO if decision else None,
            int(tracked),
            LONG_AGO if closed else None,
            jd,
            posting_id,
        ),
    )
    return posting_id


def jsa_packet(monkeypatch, capsys, *args):
    monkeypatch.setattr(sys, "argv", ["jsa", "packet", *map(str, args)])
    code = 0
    try:
        cli.main()
    except SystemExit as exit_:
        code = exit_.code if isinstance(exit_.code, int) else 1
    captured = capsys.readouterr()
    return code, captured.out + captured.err


def entries(directory: Path) -> set[str]:
    return {path.name for path in directory.iterdir()} if directory.exists() else set()


def snapshot(directory: Path) -> dict[str, bytes]:
    return {p.name: p.read_bytes() for p in directory.iterdir()}


# --- the example profile's resume ---------------------------------------------


def test_the_example_resume_opens_and_describes_the_fictional_candidate():
    document = docx.Document(str(EXAMPLE_DIR / "resume.docx"))
    text = "\n".join(p.text for p in document.paragraphs)
    assert text.strip()
    assert "Jordan Example" in text
    assert "example.com" in text


# --- the base resume ----------------------------------------------------------


@pytest.mark.parametrize("state", ["missing", "empty"])
def test_a_missing_or_empty_base_resume_fails_before_any_row_is_processed(
    pdb, env, monkeypatch, capsys, state
):
    profile, packets = env
    seed(pdb)
    if state == "missing":
        (profile / "resume.docx").unlink()
    else:
        (profile / "resume.docx").write_bytes(b"")
    code, output = jsa_packet(monkeypatch, capsys)
    assert code != 0
    assert "profile.example" in output
    assert entries(packets) == set()


def test_the_resume_requirement_applies_even_when_the_queue_is_empty(
    pdb, env, monkeypatch, capsys
):
    profile, _ = env
    (profile / "resume.docx").unlink()
    code, output = jsa_packet(monkeypatch, capsys)
    assert code != 0
    assert "profile.example" in output


# --- building a packet --------------------------------------------------------


def test_an_apply_untracked_open_posting_gets_a_folder_with_jd_and_resume_copy(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    seed(pdb)
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    folder = packets / PLAIN
    assert entries(folder) == {"job_posting.md", RESUME_COPY}
    assert (folder / "job_posting.md").read_text(encoding="utf-8") == JD
    assert (folder / RESUME_COPY).read_bytes() == (profile / "resume.docx").read_bytes()


def test_no_checklist_is_written(pdb, env, monkeypatch, capsys):
    _, packets = env
    seed(pdb)
    jsa_packet(monkeypatch, capsys)
    assert "resume_checklist.md" not in entries(packets / PLAIN)


def test_without_a_candidate_name_the_resume_copy_has_no_prefix(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    write_config_toml(profile, f'packets_dir = "{packets}"\n')
    seed(pdb)
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert entries(packets / PLAIN) == {
        "job_posting.md",
        "Resume_StaffEngineer_AcmeWidgets.docx",
    }


def test_a_null_jd_gets_a_folder_and_resume_copy_but_no_job_posting(
    pdb, env, monkeypatch, capsys
):
    _, packets = env
    seed(pdb, jd=None)
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert entries(packets / PLAIN) == {RESUME_COPY}


def test_the_folder_name_uses_the_stored_slug_and_normalized_company(
    pdb, env, monkeypatch, capsys
):
    _, packets = env
    seed(pdb, company="Foo/Bar Corp LLC", title="Senior: Data Engineer?")
    row = pdb.execute("SELECT normalized_company, title_slug FROM postings").fetchone()
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert entries(packets) == {f"{row[0]} - {row[1]}"}


def test_the_default_packets_directory_is_documents_job_applications(
    pdb, env, monkeypatch, capsys, tmp_path
):
    profile, _ = env
    write_config_toml(profile, 'candidate_name = "Pat Example"\n')
    home = tmp_path / "home"
    home.mkdir()
    monkeypatch.setenv("HOME", str(home))
    # A default left unexpanded would build under the working directory, so keep that out of the repo.
    monkeypatch.chdir(tmp_path)
    seed(pdb)
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    folder = home / "Documents" / "Job Applications" / PLAIN
    assert (folder / "job_posting.md").is_file()


# --- name collisions ----------------------------------------------------------


def test_the_lower_id_keeps_the_plain_name_and_each_higher_id_gets_its_id(
    pdb, env, monkeypatch, capsys
):
    _, packets = env
    first = seed(pdb, jd="first")
    second = seed(pdb, jd="second")
    third = seed(pdb, jd="third")
    assert first < second < third
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert entries(packets) == {PLAIN, f"{PLAIN} ({second})", f"{PLAIN} ({third})"}
    for name, jd in ((PLAIN, "first"), (f"{PLAIN} ({second})", "second")):
        assert (packets / name / "job_posting.md").read_text(encoding="utf-8") == jd


def test_the_name_rule_reads_the_database_not_the_queue(pdb, env, monkeypatch, capsys):
    """A tracked lower-id sibling still holds the plain name, so the queued one gets its id."""
    _, packets = env
    seed(pdb, tracked=True)
    second = seed(pdb)
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert entries(packets) == {f"{PLAIN} ({second})"}


def test_a_higher_id_built_first_does_not_enter_the_lower_ids_folder(
    pdb, env, monkeypatch, capsys
):
    _, packets = env
    first = seed(pdb, jd="first")
    second = seed(pdb, jd="second")
    jsa_packet(monkeypatch, capsys, "--id", second)
    assert entries(packets) == {f"{PLAIN} ({second})"}
    jsa_packet(monkeypatch, capsys, "--id", first)
    assert (packets / PLAIN / "job_posting.md").read_text(encoding="utf-8") == "first"
    other = packets / f"{PLAIN} ({second})" / "job_posting.md"
    assert other.read_text(encoding="utf-8") == "second"


# --- existing folders and files -----------------------------------------------


def test_a_posting_whose_folder_exists_is_skipped_and_nothing_changes(
    pdb, env, monkeypatch, capsys
):
    _, packets = env
    seed(pdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / "notes.txt").write_text("mine", encoding="utf-8")
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert snapshot(folder) == {"notes.txt": b"mine"}


def test_a_second_run_changes_nothing_in_the_built_folder(
    pdb, env, monkeypatch, capsys
):
    _, packets = env
    seed(pdb)
    jsa_packet(monkeypatch, capsys)
    folder = packets / PLAIN
    (folder / RESUME_COPY).write_bytes(b"my revision")
    (folder / "job_posting.md").write_text("hand edited", encoding="utf-8")
    before = snapshot(folder)
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert snapshot(folder) == before


def test_skipping_one_existing_folder_does_not_stop_the_others(
    pdb, env, monkeypatch, capsys
):
    _, packets = env
    seed(pdb, company="Beta Labs")
    seed(pdb, company="Gamma Labs")
    (packets / "Beta Labs - Staff Engineer").mkdir(parents=True)
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert entries(packets / "Gamma Labs - Staff Engineer") == {
        "job_posting.md",
        "PatExample_Resume_StaffEngineer_GammaLabs.docx",
    }
    assert entries(packets / "Beta Labs - Staff Engineer") == set()


# --- who is in the queue ------------------------------------------------------


@pytest.mark.parametrize(
    "state",
    [{"decision": "Skip"}, {"decision": None}, {"tracked": True}, {"closed": True}],
    ids=["skip", "undecided", "tracked", "closed"],
)
def test_only_apply_untracked_open_postings_are_processed(
    pdb, env, monkeypatch, capsys, state
):
    _, packets = env
    seed(pdb, **state)
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert entries(packets) == set()


@pytest.mark.parametrize("state", [{"tracked": True}, {"closed": True}])
def test_id_processes_a_tracked_or_closed_apply_posting(
    pdb, env, monkeypatch, capsys, state
):
    _, packets = env
    posting_id = seed(pdb, **state)
    code, _ = jsa_packet(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert entries(packets / PLAIN) == {"job_posting.md", RESUME_COPY}


@pytest.mark.parametrize("decision", ["Skip", None])
def test_id_never_processes_a_non_apply_posting(
    pdb, env, monkeypatch, capsys, decision
):
    _, packets = env
    posting_id = seed(pdb, decision=decision, tracked=True, closed=True)
    jsa_packet(monkeypatch, capsys, "--id", posting_id)
    assert entries(packets) == set()


def test_id_builds_only_that_posting(pdb, env, monkeypatch, capsys):
    _, packets = env
    seed(pdb, company="Beta Labs")
    wanted = seed(pdb, company="Gamma Labs")
    code, _ = jsa_packet(monkeypatch, capsys, "--id", wanted)
    assert code == 0
    assert entries(packets) == {"Gamma Labs - Staff Engineer"}


def test_an_empty_queue_exits_zero_and_creates_nothing(pdb, env, monkeypatch, capsys):
    _, packets = env
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert entries(packets) == set()


# --- dry run ------------------------------------------------------------------


def test_a_dry_run_lists_the_folders_and_creates_nothing(pdb, env, monkeypatch, capsys):
    _, packets = env
    seed(pdb, company="Beta Labs")
    second = seed(pdb, company="Beta Labs")
    seed(pdb, company="Gamma Labs")
    code, output = jsa_packet(monkeypatch, capsys, "--dry-run")
    assert code == 0
    assert str(packets / "Beta Labs - Staff Engineer") in output
    assert str(packets / f"Beta Labs - Staff Engineer ({second})") in output
    assert str(packets / "Gamma Labs - Staff Engineer") in output
    assert entries(packets) == set()


def test_a_dry_run_with_id_creates_nothing(pdb, env, monkeypatch, capsys):
    _, packets = env
    posting_id = seed(pdb)
    code, output = jsa_packet(monkeypatch, capsys, "--id", posting_id, "--dry-run")
    assert code == 0
    assert str(packets / PLAIN) in output
    assert entries(packets) == set()


def test_packet_never_writes_to_the_database(pdb, env, monkeypatch, capsys):
    """`added_to_tracker` is generate's completion guard (XC-10), so packet leaves it alone."""
    seed(pdb)
    before = pdb.execute("SELECT * FROM postings").fetchall()
    jsa_packet(monkeypatch, capsys, "--dry-run")
    jsa_packet(monkeypatch, capsys)
    assert pdb.execute("SELECT * FROM postings").fetchall() == before


# --- naming is pure -----------------------------------------------------------


@pytest.mark.parametrize(
    "candidate, expected",
    [
        ("Pat Example", "PatExample_Resume_StaffEngineer_AcmeWidgets"),
        (None, "Resume_StaffEngineer_AcmeWidgets"),
    ],
)
def test_resume_file_stem(candidate, expected):
    assert resume_file_stem(candidate, "Staff Engineer", "Acme Widgets") == expected


def test_resume_file_stem_has_no_spaces_and_no_path_hostile_characters():
    stem = resume_file_stem(
        'Pat: "P" Example',
        title_slug("Data / Analytics Lead"),
        normalize_company("A/B Corp Inc"),
    )
    assert " " not in stem
    assert not set(stem) & set('/\\:*?"<>|')


@pytest.mark.parametrize(
    "posting_id, shares, expected",
    [
        (1, False, "Acme Widgets - Staff Engineer"),
        (7, True, "Acme Widgets - Staff Engineer (7)"),
    ],
)
def test_packet_dir_name(posting_id, shares, expected):
    name = packet_dir_name(
        "Acme Widgets", "Staff Engineer", posting_id, shares_name=shares
    )
    assert name == expected


def test_naming_makes_no_file_system_database_or_network_call(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("naming reached outside itself")

    monkeypatch.setattr(socket, "socket", forbidden)
    monkeypatch.setattr(socket, "getaddrinfo", forbidden)
    monkeypatch.setattr(sqlite3, "connect", forbidden)
    monkeypatch.setattr(db, "connect", forbidden)
    monkeypatch.setattr(Path, "stat", forbidden)
    monkeypatch.setattr(Path, "exists", forbidden)
    monkeypatch.setattr(Path, "mkdir", forbidden)
    monkeypatch.setattr(Path, "open", forbidden)
    monkeypatch.setattr("builtins.open", forbidden)
    dir_name = packet_dir_name("Acme", "Staff Engineer", 3, shares_name=True)
    assert dir_name == "Acme - Staff Engineer (3)"
    assert (
        resume_file_stem("Pat", "Staff Engineer", "Acme")
        == "Pat_Resume_StaffEngineer_Acme"
    )
