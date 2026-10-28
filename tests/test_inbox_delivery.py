"""The inbox builds an emailed posting's packet and delivers it to Drive (issue #88; PRD 04 "Packets built on the inbox machine", "Tracker write"; PRD 03 outcome labels; PRD 06 `JSA_GWS_CREDENTIALS`; XC-5, XC-9, XC-10).

`gws` is replaced by a stateful fake behind `JSA_GWS_BIN` that serves Gmail, Drive, and the tracker
Sheet; Claude at the shared agent loop; outside HTTP beneath the app's client; the
database is the libSQL container; the profile is a temporary directory.
"""

import base64
import io
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

import docx
import httpx
import pytest
from conftest import REPO_ROOT, drop_all_tables
from profile_helpers import (
    ROLE_FAMILIES,
    ROLE_FAMILY,
    copy_example,
    family_dir,
    write_config_toml,
)
from test_generate import (
    JD,
    PLAIN,
    QUUX_EDIT,
    QUUX_JD,
    REDLINE_EDITS,
    RESUME_COPY,
    RESUME_PDF,
    Agent,
    column,
    failed,
    redline_reply,
    seed,
)
from test_inbox import (
    ANALYTICS_ADDRESS,
    DATA_ANALYSIS_ADDRESS,
    DESCRIPTION,
    OWNER_CREDENTIAL,
    SENDER,
    FakeGmail,
    bare_posting,
    run_inbox,
)

from jsa import agent_loop, cli, db

FOLDER_ID = "folder-123"
OTHER_FOLDER_ID = "elsewhere-456"
FOLDER_MIME = "application/vnd.google-apps.folder"
BASE_LINE = "Built the Quuxlate platform from scratch."
OTHER_FAMILY = ROLE_FAMILIES[1]
OTHER_BASE_LINE = "Reported the Zorblax metrics every week."
CONFIG = """\
candidate_name = "Pat Example"
tracker_spreadsheet_id = "sheet-123"
packets_dir = "{packets}"

[inbox]
app = "jsa-inbox-example"
senders = ["{sender}"]
drive_folder_id = "{folder}"

[agents.redline]
model = "claude-opus-5-5"
effort = "medium"
"""

# What the fake `gws` does for Gmail is the stub in test_inbox; this one adds Drive and the tracker
# Sheet, keeping the whole world (files, appended rows, every call) in $FAKE_WORLD.
FAKE_GWS = """\
#!@PYTHON@
import base64, json, os, re, sys

argv = sys.argv[1:]
tokens = argv[: argv.index("--params")] if "--params" in argv else argv
if tokens[0] not in ("drive", "sheets"):
    os.execv("@PYTHON@", ["@PYTHON@", "@GMAIL@"] + argv)

path = os.environ["FAKE_WORLD"]
state = json.load(open(path))
params = json.loads(argv[argv.index("--params") + 1])
body = json.loads(argv[argv.index("--json") + 1]) if "--json" in argv else None
upload = argv[argv.index("--upload") + 1] if "--upload" in argv else None
credential_file = os.environ.get("GOOGLE_WORKSPACE_CLI_CREDENTIALS_FILE")
credential = None
if credential_file and os.path.isfile(credential_file):
    credential = open(credential_file).read()
state["calls"].append({
    "tokens": tokens, "params": params, "body": body, "upload": upload,
    "cwd": os.getcwd(), "credential": credential, "uses_file": credential_file is not None,
})


def finish(code, out=None, err=None):
    json.dump(state, open(path, "w"))
    if err:
        print(err, file=sys.stderr)
    if out is not None:
        print(json.dumps(out))
    sys.exit(code)


method = tokens[-1]
if tokens[0] == "sheets":
    if method != "append":
        finish(1, err="unsupported sheets call")
    if state["fail_sheets"]:
        finish(1, err="500 backend error")
    state["rows"].append(body["values"][0])
    finish(0, {"updates": {"updatedRows": 1}})

if tokens[:2] != ["drive", "files"] or method not in ("list", "create"):
    state["forbidden"].append(tokens)
    finish(1, err="unsupported drive call")

QUOTED = r"'((?:\\\\.|[^'\\\\])*)'"


def unquote(text):
    return re.sub(r"\\\\(.)", r"\\1", text)


files = state["files"]
if method == "list":
    if state["fail_list"]:
        finish(1, err="500 backend error")
    query = params["q"]
    wanted = {}
    for key, pattern in (
        ("parent", QUOTED + r" in parents"),
        ("name", r"\\bname = " + QUOTED),
        ("mime", r"\\bmimeType = " + QUOTED),
    ):
        found = re.search(pattern, query)
        if found:
            wanted[key] = unquote(found.group(1))
    hits = [
        {"id": i, "name": f["name"]}
        for i, f in files.items()
        if not f["trashed"]
        and ("parent" not in wanted or wanted["parent"] in f["parents"])
        and ("name" not in wanted or f["name"] == wanted["name"])
        and ("mime" not in wanted or f["mimeType"] == wanted["mime"])
    ]
    finish(0, {"files": hits})

name = body["name"]
if name in state["fail_upload"] and upload is not None:
    finish(1, err="500 backend error")
content = None
if upload is not None:
    if os.path.isabs(upload) or ".." in upload.split("/"):
        finish(1, err="upload must be beneath the working directory")
    content = base64.b64encode(open(upload, "rb").read()).decode()
state["next"] += 1
file_id = "id%d" % state["next"]
files[file_id] = {
    "name": name, "parents": body.get("parents", []), "mimeType": body.get("mimeType"),
    "content": content, "trashed": False,
}
finish(0, {"id": file_id})
"""


class World:
    """Drive, the tracker Sheet, and the call log, as the fake `gws` keeps them."""

    def __init__(self, path: Path):
        self.path = path
        self.write(
            {
                "files": {},
                "next": 0,
                "calls": [],
                "rows": [],
                "forbidden": [],
                "fail_upload": [],
                "fail_sheets": False,
                "fail_list": False,
            }
        )

    def read(self) -> dict:
        return json.loads(self.path.read_text())

    def write(self, state: dict) -> None:
        self.path.write_text(json.dumps(state))

    def set(self, key: str, value) -> None:
        state = self.read()
        state[key] = value
        self.write(state)

    def put(self, name: str, parent: str, *, mime=None, content: bytes | None = None):
        state = self.read()
        state["next"] += 1
        file_id = f"seed{state['next']}"
        state["files"][file_id] = {
            "name": name,
            "parents": [parent],
            "mimeType": mime,
            "content": None if content is None else base64.b64encode(content).decode(),
            "trashed": False,
        }
        self.write(state)
        return file_id

    def folders(self, name: str, parent: str = FOLDER_ID) -> list[str]:
        return [
            file_id
            for file_id, f in self.read()["files"].items()
            if f["name"] == name
            and f["mimeType"] == FOLDER_MIME
            and parent in f["parents"]
        ]

    def contents(self, folder_id: str) -> dict[str, bytes]:
        found = {}
        for f in self.read()["files"].values():
            if folder_id in f["parents"] and f["mimeType"] != FOLDER_MIME:
                assert f["name"] not in found, f"{f['name']} is in the folder twice"
                found[f["name"]] = base64.b64decode(f["content"])
        return found

    def drive_calls(self, method: str | None = None) -> list[dict]:
        return [
            c
            for c in self.read()["calls"]
            if c["tokens"][0] == "drive" and method in (None, c["tokens"][-1])
        ]

    def uploads(self) -> list[dict]:
        return [c for c in self.drive_calls("create") if c["upload"] is not None]

    def sheet_calls(self) -> list[dict]:
        return [c for c in self.read()["calls"] if c["tokens"][0] == "sheets"]

    def rows(self) -> list[list]:
        return self.read()["rows"]


@pytest.fixture
def conn(db_url):
    """A connection to a database holding no postings but the test's own: a packet's folder name depends on its namesakes."""
    drop_all_tables(db_url)
    connection = db.connect()
    yield connection
    connection.close()


class Requests:
    """Outside HTTP: every request is logged and answered 404, so any fetch of a posting reads as closed."""

    def __init__(self):
        self.seen = []


@pytest.fixture
def requests(monkeypatch):
    log = Requests()

    def handle(_transport, request):
        log.seen.append(str(request.url))
        return httpx.Response(404, json={"error": "not found"})

    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", handle)
    return log


@pytest.fixture
def agent(monkeypatch):
    stand_in = Agent()
    monkeypatch.setattr(agent_loop, "query", stand_in.query)
    return stand_in


@pytest.fixture
def env(tmp_path, monkeypatch):
    """A profile with an `[inbox]` table and a distinctive base resume per role family; returns (profile, packets_dir)."""
    from test_inbox import FAKE_GWS as GMAIL_STUB

    profile = copy_example(tmp_path / "profile")
    packets = tmp_path / "packets"
    write_config_toml(
        profile, CONFIG.format(packets=packets, sender=SENDER, folder=FOLDER_ID)
    )
    for family, line in ((ROLE_FAMILY, BASE_LINE), (OTHER_FAMILY, OTHER_BASE_LINE)):
        document = docx.Document()
        document.add_paragraph().add_run("Riley Resumeperson").bold = True
        document.add_paragraph(line)
        document.save(family_dir(profile, family) / "resume.docx")
    monkeypatch.setenv("JSA_PROFILE_DIR", str(profile))
    gmail_stub = tmp_path / "gmail-gws"
    gmail_stub.write_text(GMAIL_STUB.format(python=sys.executable), encoding="utf-8")
    gmail_stub.chmod(0o755)
    stub = tmp_path / "fake-gws"
    stub.write_text(
        FAKE_GWS.replace("@PYTHON@", sys.executable).replace(
            "@GMAIL@", str(gmail_stub)
        ),
        encoding="utf-8",
    )
    stub.chmod(0o755)
    monkeypatch.setenv("JSA_GWS_BIN", str(stub))
    monkeypatch.setenv("FAKE_MAILBOX", str(tmp_path / "mailbox.json"))
    monkeypatch.setenv("FAKE_WORLD", str(tmp_path / "world.json"))
    monkeypatch.setenv("JSA_INBOX_GWS_CREDENTIALS", "MAILBOX-CREDENTIAL-111")
    monkeypatch.setenv("JSA_GWS_CREDENTIALS", OWNER_CREDENTIAL)
    monkeypatch.delenv("GOOGLE_WORKSPACE_CLI_CREDENTIALS_FILE", raising=False)
    monkeypatch.delenv("JSA_GENERATE_WORKERS", raising=False)
    return profile, packets


@pytest.fixture
def gmail(env, tmp_path) -> FakeGmail:
    return FakeGmail(tmp_path / "mailbox.json")


@pytest.fixture
def world(env, tmp_path) -> World:
    return World(tmp_path / "world.json")


def give_role_family(conn, posting_id, role_family) -> None:
    conn.execute(
        "UPDATE postings SET role_family = ? WHERE id = ?", (role_family, posting_id)
    )


def email(gmail, conn, message_id="m1", role_family=ROLE_FAMILY, **seed_fields) -> int:
    """Store an Apply posting that has a description and a role family, and email its URL; returns its id."""
    posting_id = seed(conn, **seed_fields)
    give_role_family(conn, posting_id, role_family)
    url = conn.execute(
        "SELECT url FROM postings WHERE id = ?", (posting_id,)
    ).fetchone()[0]
    gmail.add(message_id, 1000 + len(message_id), f"{url}\n")
    return posting_id


def tracked(conn, posting_id) -> bool:
    return bool(column(conn, posting_id, "added_to_tracker"))


# --- acceptance 1 + 2 + 3: build, upload, then track ---------------------------


def test_an_emailed_posting_with_a_description_ends_tracked_with_its_packet_in_drive(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    agent.respond = redline_reply(QUUX_EDIT)
    posting_id = email(gmail, conn, jd=QUUX_JD)
    assert run_inbox(monkeypatch) == 0
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    assert not gmail.in_inbox("m1")
    assert tracked(conn, posting_id)
    [row] = world.rows()
    assert row[0] == posting_id
    [folder_id] = world.folders(PLAIN)
    uploaded = world.contents(folder_id)
    # The redline edited the copy, so the packet has one resume file and no resume PDF.
    assert set(uploaded) == {RESUME_COPY, "job_posting.md", REDLINE_EDITS}
    assert uploaded["job_posting.md"].decode() == QUUX_JD


def test_the_redlined_resume_copy_is_uploaded_as_the_exact_docx_file(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    profile, _ = env
    agent.respond = redline_reply(QUUX_EDIT)
    email(gmail, conn, jd=QUUX_JD)
    run_inbox(monkeypatch)
    [folder_id] = world.folders(PLAIN)
    uploaded = world.contents(folder_id)
    assert uploaded[RESUME_COPY] != (family_dir(profile) / "resume.docx").read_bytes()
    with zipfile.ZipFile(io.BytesIO(uploaded[RESUME_COPY])) as archive:
        body = archive.read("word/document.xml").decode()
        comments = archive.read("word/comments.xml").decode()
    assert "<w:ins " in body and "<w:del " in body
    assert "REASON-MARKER" in comments


def test_a_resume_the_redline_left_alone_is_delivered_with_its_pdf(
    db_url, conn, env, gmail, world, agent, soffice, requests, monkeypatch
):
    profile, _ = env
    email(gmail, conn, jd=QUUX_JD)
    assert run_inbox(monkeypatch) == 0
    [folder_id] = world.folders(PLAIN)
    uploaded = world.contents(folder_id)
    assert uploaded[RESUME_COPY] == (family_dir(profile) / "resume.docx").read_bytes()
    assert uploaded[RESUME_PDF].decode() == f"PDF OF: {RESUME_COPY}"


def test_a_docx_cover_letter_is_delivered_with_its_pdf(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    profile, _ = env
    (family_dir(profile) / "cover_letter.docx").write_bytes(b"THE COVER LETTER")
    agent.respond = redline_reply(QUUX_EDIT)
    email(gmail, conn, jd=QUUX_JD)
    assert run_inbox(monkeypatch) == 0
    [folder_id] = world.folders(PLAIN)
    uploaded = world.contents(folder_id)
    letter = RESUME_COPY.replace("_Resume_", "_CoverLetter_")
    assert uploaded[letter] == b"THE COVER LETTER"
    assert uploaded[letter.replace(".docx", ".pdf")].decode() == f"PDF OF: {letter}"


def test_a_cover_letter_in_the_profile_is_delivered_with_the_packet(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    profile, _ = env
    (family_dir(profile) / "cover_letter.pdf").write_bytes(b"THE COVER LETTER")
    agent.respond = redline_reply(QUUX_EDIT)
    email(gmail, conn, jd=QUUX_JD)
    assert run_inbox(monkeypatch) == 0
    [folder_id] = world.folders(PLAIN)
    uploaded = world.contents(folder_id)
    letter = RESUME_COPY.replace("_Resume_", "_CoverLetter_").replace(".docx", ".pdf")
    assert uploaded[letter] == b"THE COVER LETTER"


def test_without_a_cover_letter_the_delivered_packet_has_none(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    agent.respond = redline_reply(QUUX_EDIT)
    email(gmail, conn, jd=QUUX_JD)
    assert run_inbox(monkeypatch) == 0
    [folder_id] = world.folders(PLAIN)
    assert not [name for name in world.contents(folder_id) if "CoverLetter" in name]


def test_uploads_go_into_the_configured_folder_and_are_never_converted(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    email(gmail, conn)
    run_inbox(monkeypatch)
    state = world.read()
    [folder_id] = world.folders(PLAIN)
    assert state["files"][folder_id]["parents"] == [FOLDER_ID]
    inside = [f for f in state["files"].values() if folder_id in f["parents"]]
    assert inside
    for f in inside:
        assert not (f["mimeType"] or "").startswith("application/vnd.google-apps")
    # Only the packet folder itself was created at the top.
    assert [
        f["name"] for f in state["files"].values() if FOLDER_ID in f["parents"]
    ] == [PLAIN]


def test_every_file_is_uploaded_from_beneath_the_packet_folder_with_the_owners_credential(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    email(gmail, conn)
    run_inbox(monkeypatch)
    uploads = world.uploads()
    assert uploads
    for call in uploads:
        relative = call["upload"]
        assert not Path(relative).is_absolute() and ".." not in Path(relative).parts
        assert Path(call["cwd"]).name == PLAIN
        assert relative == call["body"]["name"]
    for call in world.drive_calls() + world.sheet_calls():
        assert call["credential"] == OWNER_CREDENTIAL


def test_the_tracker_row_is_appended_only_after_every_file_is_uploaded(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    email(gmail, conn)
    run_inbox(monkeypatch)
    calls = world.read()["calls"]
    kinds = [c["tokens"][0] for c in calls]
    assert kinds.count("sheets") == 1
    first_append = kinds.index("sheets")
    assert "drive" not in kinds[first_append:]
    uploaded = {c["upload"] for c in world.uploads()}
    assert {RESUME_COPY, "job_posting.md", REDLINE_EDITS} <= uploaded


def test_the_packet_is_built_in_a_temporary_directory_not_in_packets_dir(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    _, packets = env
    email(gmail, conn)
    run_inbox(monkeypatch)
    assert not packets.exists() or not list(packets.iterdir())
    cwds = {c["cwd"] for c in world.uploads()}
    assert cwds
    assert all(not Path(cwd).is_relative_to(packets) for cwd in cwds)


def test_the_build_makes_one_model_call_the_redline(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    email(gmail, conn)
    run_inbox(monkeypatch)
    assert len(agent.calls) == 1
    assert JD.strip() in agent.prompts[0]
    assert BASE_LINE in agent.prompts[0]


def test_the_redline_on_the_inbox_machine_reads_the_skills_list_from_the_image(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    """The machine's profile holds the image's `search/` and skills list plus the machine files, and nothing else."""
    profile, _ = env
    (profile / "resume_skills.md").write_text(
        "INBOX-SKILL-ALPHA\nINBOX-SKILL-BETA\n", encoding="utf-8"
    )
    machine = {"search", "resume_skills.md", "config.toml", "role_families"}
    for entry in profile.iterdir():
        if entry.name not in machine:
            shutil.rmtree(entry) if entry.is_dir() else entry.unlink()
    email(gmail, conn)
    run_inbox(monkeypatch)
    (prompt,) = agent.prompts
    assert {"INBOX-SKILL-ALPHA", "INBOX-SKILL-BETA"} <= set(prompt.splitlines())


def test_the_delivered_packet_is_what_jsa_generate_id_builds_locally(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch, capsys
):
    _, packets = env
    agent.respond = redline_reply(QUUX_EDIT)
    posting_id = email(gmail, conn, jd=QUUX_JD)
    run_inbox(monkeypatch)
    [folder_id] = world.folders(PLAIN)
    uploaded = world.contents(folder_id)
    monkeypatch.setattr(sys, "argv", ["jsa", "generate", "--id", str(posting_id)])
    try:
        cli.main()
    except SystemExit as exit_:
        assert not exit_.code
    local = packets / PLAIN
    assert {p.name for p in local.iterdir()} == set(uploaded)
    for path in local.iterdir():
        if path.suffix != ".docx":
            assert path.read_bytes() == uploaded[path.name], path.name


def test_the_packet_is_built_from_the_postings_role_family_and_tracked_on_its_tab(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    profile, _ = env
    other_cover_letter = family_dir(profile, OTHER_FAMILY) / "cover_letter.docx"
    other_cover_letter.write_bytes(b"THE OTHER COVER LETTER")
    family_dir(profile).joinpath("cover_letter.docx").write_bytes(b"THE FIRST LETTER")
    posting_id = email(gmail, conn, role_family=OTHER_FAMILY)
    assert run_inbox(monkeypatch) == 0
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    assert tracked(conn, posting_id)
    assert OTHER_BASE_LINE in agent.prompts[0]
    assert BASE_LINE not in agent.prompts[0]
    [folder_id] = world.folders(PLAIN)
    uploaded = world.contents(folder_id)
    letter = RESUME_COPY.replace("_Resume_", "_CoverLetter_")
    assert uploaded[letter] == b"THE OTHER COVER LETTER"
    assert set(uploaded) >= {RESUME_COPY, letter}
    [call] = world.sheet_calls()
    assert call["params"]["range"] == f"'{OTHER_FAMILY}'!A:H"


@pytest.mark.parametrize(
    ("role_family", "label"),
    [(None, "jsa/needs-family"), ("Retired Family", "jsa/failed")],
    ids=["none", "unknown"],
)
def test_a_posting_without_a_usable_role_family_is_labeled_and_stays_untracked(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch, role_family, label
):
    """PRD 03: no role family (plain address, several families) is `needs-family`; one the profile lacks is `failed`."""
    posting_id = email(gmail, conn, role_family=role_family)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == [label]
    assert not gmail.in_inbox("m1")
    assert agent.calls == []
    assert world.drive_calls("create") == [] and world.sheet_calls() == []
    assert not tracked(conn, posting_id)
    assert column(conn, posting_id, "decision") == "Apply"
    assert column(conn, posting_id, "role_family") == role_family


# --- the address names the role family (issue #196; PRD 03, PRD 04) -------------


def test_a_tagged_message_with_a_description_builds_from_that_role_family_and_tracks_on_its_tab(
    db_url, conn, env, gmail, world, agent, web, monkeypatch
):
    url = bare_posting(web)
    gmail.add("m1", 1000, f"{url}\n\n{DESCRIPTION}\n", to=DATA_ANALYSIS_ADDRESS)
    assert run_inbox(monkeypatch) == 0
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    row = conn.execute(
        "SELECT id, role_family, added_to_tracker FROM postings WHERE url = ?", (url,)
    ).fetchone()
    assert row[1:] == (OTHER_FAMILY, 1)
    assert OTHER_BASE_LINE in agent.prompts[0]
    assert BASE_LINE not in agent.prompts[0]
    [call] = world.sheet_calls()
    assert call["params"]["range"] == f"'{OTHER_FAMILY}'!A:H"
    [appended] = world.rows()
    assert appended[0] == row[0]


def test_the_other_roles_tag_builds_from_the_other_role_family(
    db_url, conn, env, gmail, world, agent, web, monkeypatch
):
    url = bare_posting(web)
    gmail.add("m1", 1000, f"{url}\n\n{DESCRIPTION}\n", to=ANALYTICS_ADDRESS)
    assert run_inbox(monkeypatch) == 0
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    assert BASE_LINE in agent.prompts[0]
    assert OTHER_BASE_LINE not in agent.prompts[0]
    [call] = world.sheet_calls()
    assert call["params"]["range"] == f"'{ROLE_FAMILY}'!A:H"


@pytest.mark.parametrize(
    "address",
    ["jobs.mailbox@example.net", "jobs.mailbox+no-such-family@example.net"],
    ids=["plain-address", "unmatched-tag"],
)
def test_with_two_role_families_no_matching_tag_builds_nothing_even_with_a_description(
    db_url, conn, env, gmail, world, agent, web, monkeypatch, address
):
    url = bare_posting(web)
    gmail.add("m1", 1000, f"{url}\n\n{DESCRIPTION}\n", to=address)
    assert run_inbox(monkeypatch) == 0
    assert gmail.outcome_labels("m1") == ["jsa/needs-family"]
    assert agent.calls == []
    assert world.drive_calls("create") == [] and world.sheet_calls() == []
    row = conn.execute(
        "SELECT role_family, decision, added_to_tracker, jd_markdown FROM postings WHERE url = ?",
        (url,),
    ).fetchone()
    assert row == (None, "Apply", 0, DESCRIPTION)


@pytest.mark.parametrize("name", ROLE_FAMILIES)
def test_with_one_role_family_the_plain_address_builds_from_it(
    db_url, conn, env, gmail, world, agent, web, monkeypatch, name
):
    profile, _ = env
    shutil.rmtree(family_dir(profile, next(f for f in ROLE_FAMILIES if f != name)))
    url = bare_posting(web)
    gmail.add("m1", 1000, f"{url}\n\n{DESCRIPTION}\n")
    assert run_inbox(monkeypatch) == 0
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    [call] = world.sheet_calls()
    assert call["params"]["range"] == f"'{name}'!A:H"


def test_a_tag_does_not_change_the_role_family_an_existing_row_is_built_from(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    posting_id = email(gmail, conn, role_family=ROLE_FAMILY)
    url = conn.execute(
        "SELECT url FROM postings WHERE id = ?", (posting_id,)
    ).fetchone()[0]
    gmail.add("m2", 2000, f"{url}\n", to=DATA_ANALYSIS_ADDRESS)
    assert run_inbox(monkeypatch) == 0
    assert column(conn, posting_id, "role_family") == ROLE_FAMILY
    assert gmail.outcome_labels("m2") == ["jsa/tracked"]
    assert BASE_LINE in agent.prompts[0]
    assert OTHER_BASE_LINE not in agent.prompts[0]


# --- acceptance 1: no liveness re-check ----------------------------------------


def test_a_posting_whose_page_has_closed_is_still_built_delivered_and_tracked(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    posting_id = email(gmail, conn)
    run_inbox(monkeypatch)
    assert requests.seen == []
    assert column(conn, posting_id, "closed_at") is None
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    assert world.folders(PLAIN)
    assert tracked(conn, posting_id)


def test_an_emailed_description_goes_all_the_way_to_a_delivered_packet(
    db_url, conn, env, gmail, world, agent, web, monkeypatch
):
    url = bare_posting(web)
    # The address-tag rule that gives a new emailed row its role family is PRD 03's; here the row has one already.
    db.insert_posting(
        conn,
        company="Example",
        title="Staff Platform Engineer",
        url=url,
        search_agent="manual",
        role_family=ROLE_FAMILY,
    )
    gmail.add("m1", 1000, f"{url}\n\n{DESCRIPTION}\n")
    assert run_inbox(monkeypatch) == 0
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    row = conn.execute(
        "SELECT id, added_to_tracker, jd_markdown FROM postings WHERE url = ?", (url,)
    ).fetchone()
    assert row[1] == 1 and row[2] == DESCRIPTION
    [appended] = world.rows()
    assert appended[0] == row[0]
    [folder_id] = world.folders("Example - Staff Platform Engineer")
    assert "job_posting.md" in world.contents(folder_id)


@pytest.fixture
def web(requests, monkeypatch):
    """Route outside HTTP to `web[(host, path)]`; anything unrouted is a 404 (as test_inbox's `web`)."""
    routes = {}

    def handle(_transport, request):
        requests.seen.append(str(request.url))
        answer = routes.get((request.url.host, request.url.path))
        if answer is None:
            return httpx.Response(404, json={"error": "not found"})
        if isinstance(answer, httpx.Response):
            return answer
        return httpx.Response(200, json=answer)

    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", handle)
    return routes


# --- acceptance 4: a retry reuses the folder and uploads only what is missing ---


def test_a_failed_upload_leaves_the_posting_untracked_and_a_retry_uploads_only_the_missing_file(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    posting_id = email(gmail, conn)
    world.set("fail_upload", [REDLINE_EDITS])
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/failed"]
    assert not gmail.in_inbox("m1")
    assert not tracked(conn, posting_id)
    assert world.rows() == [] and world.sheet_calls() == []
    [folder_id] = world.folders(PLAIN)
    first = world.contents(folder_id)
    assert REDLINE_EDITS not in first and "job_posting.md" in first
    uploads_before = len(world.uploads())

    world.set("fail_upload", [])
    gmail.requeue("m1")
    run_inbox(monkeypatch)
    assert "jsa/tracked" in gmail.outcome_labels("m1")
    assert tracked(conn, posting_id)
    assert world.folders(PLAIN) == [folder_id]
    after = world.contents(folder_id)
    assert set(after) == set(first) | {REDLINE_EDITS}
    assert [c["upload"] for c in world.uploads()[uploads_before:]] == [REDLINE_EDITS]
    assert len(world.rows()) == 1


def test_a_failed_tracker_append_is_labeled_failed_and_the_retry_uploads_nothing_new(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    posting_id = email(gmail, conn)
    world.set("fail_sheets", True)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/failed"]
    assert not tracked(conn, posting_id)
    assert world.rows() == []
    [folder_id] = world.folders(PLAIN)
    uploads_before = len(world.uploads())
    assert uploads_before > 0

    world.set("fail_sheets", False)
    gmail.requeue("m1")
    run_inbox(monkeypatch)
    assert tracked(conn, posting_id)
    assert len(world.uploads()) == uploads_before
    assert world.folders(PLAIN) == [folder_id]
    assert len(world.rows()) == 1
    assert "jsa/tracked" in gmail.outcome_labels("m1")


def test_a_retry_never_replaces_or_deletes_a_file_already_in_the_folder(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    folder_id = world.put(PLAIN, FOLDER_ID, mime=FOLDER_MIME)
    world.put(RESUME_COPY, folder_id, content=b"THE USER'S OWN EDITS")
    posting_id = email(gmail, conn)
    run_inbox(monkeypatch)
    assert world.folders(PLAIN) == [folder_id]
    contents = world.contents(folder_id)
    assert contents[RESUME_COPY] == b"THE USER'S OWN EDITS"
    assert {"job_posting.md", REDLINE_EDITS} <= set(contents)
    assert RESUME_COPY not in {c["upload"] for c in world.uploads()}
    state = world.read()
    assert state["forbidden"] == []
    assert {c["tokens"][-1] for c in world.drive_calls()} <= {"list", "create"}
    assert all(not f["trashed"] for f in state["files"].values())
    assert tracked(conn, posting_id)


def test_a_folder_of_the_same_name_elsewhere_is_not_reused(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    stranger = world.put(PLAIN, OTHER_FOLDER_ID, mime=FOLDER_MIME)
    world.put("job_posting.md", stranger, content=b"NOT OURS")
    email(gmail, conn)
    run_inbox(monkeypatch)
    [ours] = world.folders(PLAIN)
    assert ours != stranger
    assert world.contents(stranger) == {"job_posting.md": b"NOT OURS"}
    assert "job_posting.md" in world.contents(ours)


def test_a_trashed_packet_folder_is_not_reused(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    trashed = world.put(PLAIN, FOLDER_ID, mime=FOLDER_MIME)
    state = world.read()
    state["files"][trashed]["trashed"] = True
    world.write(state)
    email(gmail, conn)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    live = [
        f
        for f in world.read()["files"].values()
        if f["name"] == PLAIN and not f["trashed"]
    ]
    assert len(live) == 1


def test_a_packet_folder_name_with_a_quote_in_it_is_found_again_on_the_retry(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    posting_id = email(gmail, conn, company="O'Brien Labs", title="Staff Engineer")
    world.set("fail_sheets", True)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/failed"]
    world.set("fail_sheets", False)
    gmail.requeue("m1")
    run_inbox(monkeypatch)
    assert tracked(conn, posting_id)
    names = [
        f["name"]
        for f in world.read()["files"].values()
        if f["mimeType"] == FOLDER_MIME
    ]
    assert len(names) == 1


# --- acceptance 5: tracked postings are not rebuilt ----------------------------


def test_an_already_tracked_posting_is_neither_rebuilt_nor_uploaded_and_is_labeled_tracked(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    posting_id = email(gmail, conn, tracked=True)
    assert run_inbox(monkeypatch) == 0
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    assert not gmail.in_inbox("m1")
    assert agent.calls == []
    assert world.drive_calls() == [] and world.sheet_calls() == []
    assert world.rows() == []
    assert requests.seen == []
    assert tracked(conn, posting_id)


def test_a_posting_tracked_by_an_earlier_message_is_not_delivered_again(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    email(gmail, conn)
    url = conn.execute("SELECT url FROM postings ORDER BY id DESC LIMIT 1").fetchone()[
        0
    ]
    run_inbox(monkeypatch)
    uploads = len(world.uploads())
    builds = len(agent.calls)
    gmail.add("m2", 5000, f"{url}\n")
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m2") == ["jsa/tracked"]
    assert len(world.uploads()) == uploads
    assert len(agent.calls) == builds
    assert len(world.rows()) == 1


# --- acceptance 6: failures are labeled failed and leave the posting untracked ---


def test_a_failed_pdf_render_is_labeled_failed_and_nothing_is_delivered(
    db_url, conn, env, gmail, world, agent, soffice, requests, monkeypatch
):
    monkeypatch.setenv("SOFFICE_FAIL", "1")
    posting_id = email(gmail, conn)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/failed"]
    assert not tracked(conn, posting_id)
    assert world.drive_calls("create") == [] and world.sheet_calls() == []


def test_a_failed_redline_is_labeled_failed_and_nothing_is_delivered(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    agent.respond = lambda prompt: failed(result="Claude API error")
    posting_id = email(gmail, conn)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/failed"]
    assert not tracked(conn, posting_id)
    assert world.drive_calls("create") == [] and world.sheet_calls() == []


def test_an_unreadable_drive_folder_listing_is_labeled_failed_and_tracks_nothing(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    world.set("fail_list", True)
    posting_id = email(gmail, conn)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/failed"]
    assert not tracked(conn, posting_id)
    assert world.rows() == []


def test_one_failed_posting_does_not_stop_the_next_message(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    first = email(gmail, conn, "m1", company="First Labs")
    second = email(gmail, conn, "m2", company="Second Labs")
    world.set("fail_upload", ["job_posting.md"])
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/failed"]
    assert gmail.outcome_labels("m2") == ["jsa/failed"]
    world.set("fail_upload", [])
    gmail.requeue("m1")
    gmail.requeue("m2")
    run_inbox(monkeypatch)
    # Opened after the runs: a connection idle through them could have its stream expire.
    after = db.connect()
    assert tracked(after, first) and tracked(after, second)
    after.close()
    assert sorted(r[0] for r in world.rows()) == sorted([first, second])


def test_a_posting_without_a_description_builds_and_uploads_nothing(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    posting_id = email(gmail, conn, jd=None)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/needs-jd"]
    assert agent.calls == []
    assert world.drive_calls() == [] and world.sheet_calls() == []
    assert not tracked(conn, posting_id)


def test_a_missing_drive_folder_setting_fails_the_message_before_any_build(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    profile, packets = env
    write_config_toml(
        profile,
        CONFIG.format(packets=packets, sender=SENDER, folder="").replace(
            'drive_folder_id = ""\n', ""
        ),
    )
    posting_id = email(gmail, conn)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/failed"]
    assert agent.calls == []
    assert world.drive_calls() == [] and world.rows() == []
    assert not tracked(conn, posting_id)


# --- the owner's credential (PRD 06; convention 3) -----------------------------


def test_a_tracker_append_outside_the_inbox_uses_the_owners_credential_when_one_is_set(
    db_url, conn, env, world, monkeypatch, capsys
):
    posting_id = seed(conn)
    give_role_family(conn, posting_id, ROLE_FAMILY)
    monkeypatch.setattr(sys, "argv", ["jsa", "track", "--id", str(posting_id)])
    try:
        cli.main()
    except SystemExit as exit_:
        assert not exit_.code
    [call] = world.sheet_calls()
    assert call["credential"] == OWNER_CREDENTIAL
    assert tracked(conn, posting_id)


def test_with_no_owner_credential_gws_keeps_its_own_login(
    db_url, conn, env, world, monkeypatch, capsys
):
    monkeypatch.delenv("JSA_GWS_CREDENTIALS")
    posting_id = seed(conn)
    give_role_family(conn, posting_id, ROLE_FAMILY)
    monkeypatch.setattr(sys, "argv", ["jsa", "track", "--id", str(posting_id)])
    try:
        cli.main()
    except SystemExit as exit_:
        assert not exit_.code
    [call] = world.sheet_calls()
    assert call["uses_file"] is False and call["credential"] is None
    assert tracked(conn, posting_id)


def test_with_no_owner_credential_the_inbox_delivery_runs_on_gws_own_login(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    monkeypatch.delenv("JSA_GWS_CREDENTIALS")
    posting_id = email(gmail, conn)
    run_inbox(monkeypatch)
    assert gmail.outcome_labels("m1") == ["jsa/tracked"]
    assert tracked(conn, posting_id)
    for call in world.drive_calls() + world.sheet_calls():
        assert call["credential"] is None and call["uses_file"] is False


def test_the_owner_credential_is_never_used_for_gmail(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch
):
    email(gmail, conn)
    run_inbox(monkeypatch)
    gmail_calls = gmail.calls()
    assert gmail_calls
    assert all(call["credentials"] == ["mailbox"] for call in gmail_calls)


def test_the_owner_credential_text_reaches_no_log_line(
    db_url, conn, env, gmail, world, agent, requests, monkeypatch, caplog
):
    email(gmail, conn)
    run_inbox(monkeypatch)
    assert OWNER_CREDENTIAL not in caplog.text


# --- acceptance 8: .env.example -------------------------------------------------


def test_env_example_lists_the_owner_credential_as_inbox_app_only():
    lines = (REPO_ROOT / ".env.example").read_text(encoding="utf-8").splitlines()
    [index] = [
        i for i, l in enumerate(lines) if re.match(r"^#?\s*JSA_GWS_CREDENTIALS=", l)
    ]
    assert lines[index].lstrip("# ").strip() == "JSA_GWS_CREDENTIALS="
    block = []
    for line in reversed(lines[:index]):
        if not line.strip():
            break
        block.append(line)
    assert block, "JSA_GWS_CREDENTIALS has no explanatory comment"
    assert "inbox" in " ".join(block).lower()
