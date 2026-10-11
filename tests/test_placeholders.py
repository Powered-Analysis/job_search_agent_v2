"""Placeholders in the resume and cover letter (issue #130; PRD 04 "Placeholders in the resume and cover letter" and "Packet directory"; XC-9, XC-11).

`fill_placeholders` is exercised directly (it is pure over a loaded document), and the copy step through
`jsa packet`. The database is the libSQL container (and a `file:` database); the profile and the packets
directory are temporary directories.
"""

import copy
import hashlib
import io
from datetime import date, datetime
from zoneinfo import ZoneInfo

import docx
import pytest
from conftest import drop_all_tables
from docx.oxml import parse_xml
from profile_helpers import (
    SEARCH_TOML,
    copy_example,
    write_config_toml,
    write_search_toml,
)
from test_packet import JD, jsa_packet, seed

from jsa import db
from jsa.placeholders import fill_placeholders

COMPANY = "Acme Widgets"
TITLE = "Staff Engineer"
TODAY = date(2026, 10, 10)
TEXT_BOX = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    "<w:pict><w:txbxContent><w:p><w:r><w:t>[COMPANY] in a box</w:t></w:r></w:p>"
    "</w:txbxContent></w:pict></w:r>"
)
MONTHS = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]


def fill(document, company=COMPANY, title=TITLE, today=TODAY):
    return fill_placeholders(document, company, title, today)


def paragraph_with_runs(document, *texts, bold_first=False):
    paragraph = document.add_paragraph()
    for index, text in enumerate(texts):
        run = paragraph.add_run(text)
        run.bold = True if bold_first and index == 0 else None
    return paragraph


def reloaded(document):
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return docx.Document(buffer)


def body_texts(document):
    return [paragraph.text for paragraph in document.paragraphs]


# --- the filling (pure, XC-9) -------------------------------------------------


def test_each_token_is_replaced_with_its_value():
    document = docx.Document()
    document.add_paragraph("Dear [COMPANY], re [TITLE], on [DATE].")
    assert fill(document) is True
    assert body_texts(document) == [
        "Dear Acme Widgets, re Staff Engineer, on October 10, 2026."
    ]


def test_a_token_repeated_in_one_paragraph_is_filled_every_time():
    document = docx.Document()
    document.add_paragraph("[COMPANY] [COMPANY] [TITLE][TITLE]")
    fill(document)
    assert body_texts(document) == [
        "Acme Widgets Acme Widgets Staff EngineerStaff Engineer"
    ]


@pytest.mark.parametrize(
    ("day", "expected"),
    [
        (date(2026, 10, 10), "October 10, 2026"),
        (date(2026, 3, 5), "March 5, 2026"),
        (date(2027, 1, 1), "January 1, 2027"),
        (date(2028, 2, 29), "February 29, 2028"),
        (date(2026, 12, 31), "December 31, 2026"),
    ],
)
def test_date_is_the_full_month_name_the_day_without_a_leading_zero_and_the_year(
    day, expected
):
    document = docx.Document()
    document.add_paragraph("[DATE]")
    fill(document, today=day)
    assert body_texts(document) == [expected]


def test_date_uses_the_english_month_name_for_every_month():
    for month, name in enumerate(MONTHS, start=1):
        document = docx.Document()
        document.add_paragraph("[DATE]")
        fill(document, today=date(2026, month, 15))
        assert body_texts(document) == [f"{name} 15, 2026"]


def test_tokens_are_filled_in_tables_headers_and_footers():
    document = docx.Document()
    document.add_paragraph("body [COMPANY]")
    cell = document.add_table(rows=1, cols=2).rows[0].cells
    cell[0].text = "cell [TITLE]"
    nested = cell[1].add_table(rows=1, cols=1)
    nested.rows[0].cells[0].text = "nested [DATE]"
    section = document.sections[0]
    section.header.paragraphs[0].text = "head [COMPANY]"
    section.footer.paragraphs[0].text = "foot [TITLE]"
    assert fill(document) is True
    assert document.paragraphs[0].text == "body Acme Widgets"
    assert document.tables[0].rows[0].cells[0].text == "cell Staff Engineer"
    assert (
        document.tables[0].rows[0].cells[1].tables[0].rows[0].cells[0].text
        == "nested October 10, 2026"
    )
    assert section.header.paragraphs[0].text == "head Acme Widgets"
    assert section.footer.paragraphs[0].text == "foot Staff Engineer"


def test_the_first_page_and_even_page_headers_and_footers_are_filled():
    document = docx.Document()
    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "first [COMPANY]"
    section.first_page_footer.paragraphs[0].text = "first foot [TITLE]"
    section.even_page_header.paragraphs[0].text = "even [COMPANY]"
    section.even_page_footer.paragraphs[0].text = "even foot [DATE]"
    fill(document)
    assert section.first_page_header.paragraphs[0].text == "first Acme Widgets"
    assert section.first_page_footer.paragraphs[0].text == "first foot Staff Engineer"
    assert section.even_page_header.paragraphs[0].text == "even Acme Widgets"
    assert section.even_page_footer.paragraphs[0].text == "even foot October 10, 2026"


def test_a_later_sections_own_header_is_filled():
    document = docx.Document()
    document.add_paragraph("one")
    document.add_section()
    second = document.sections[1]
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "second [COMPANY]"
    fill(document)
    assert second.header.paragraphs[0].text == "second Acme Widgets"


def test_a_document_with_only_a_token_in_a_header_reports_it_held_one():
    document = docx.Document()
    document.add_paragraph("nothing here")
    document.sections[0].header.paragraphs[0].text = "[TITLE]"
    assert fill(document) is True


def test_text_boxes_are_not_filled():
    document = docx.Document()
    paragraph = document.add_paragraph("keep")
    paragraph._p.append(parse_xml(TEXT_BOX))
    before = paragraph._p.xml
    assert fill(document) is False
    assert paragraph._p.xml == before


def test_a_token_in_a_text_box_does_not_stop_the_body_being_filled():
    document = docx.Document()
    document.add_paragraph("[COMPANY]")._p.append(parse_xml(TEXT_BOX))
    fill(document)
    xml = document.element.body.xml
    assert "Acme Widgets" in xml
    assert "[COMPANY] in a box" in xml


def test_a_token_split_across_runs_is_filled_in_the_first_characters_formatting():
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Apply to ")
    paragraph.add_run("[COM").bold = True
    paragraph.add_run("PA").italic = True
    paragraph.add_run("NY]").underline = True
    paragraph.add_run(" today")
    assert fill(document) is True
    assert paragraph.text == "Apply to Acme Widgets today"
    [holder] = [run for run in paragraph.runs if "Acme Widgets" in run.text]
    assert holder.bold is True
    assert not holder.italic and not holder.underline


def test_a_token_split_one_character_per_run_is_filled():
    document = docx.Document()
    paragraph = paragraph_with_runs(document, *"[TITLE]", "!")
    fill(document)
    assert paragraph.text == "Staff Engineer!"


def test_every_other_paragraph_and_run_is_unchanged():
    document = docx.Document()
    first = paragraph_with_runs(document, "Plain ", "bold", bold_first=False)
    first.runs[1].bold = True
    target = paragraph_with_runs(document, "to [COM", "PANY] ok", bold_first=True)
    last = document.add_paragraph("Last [Company] line")
    before = [copy.deepcopy(p._p.xml) for p in (first, last)]
    fill(document)
    assert [first._p.xml, last._p.xml] == before
    assert target.text == "to Acme Widgets ok"
    assert target.runs[0].bold is True


def test_text_around_a_token_keeps_its_own_runs_and_formatting():
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Hello ").bold = True
    paragraph.add_run("[TITLE]").italic = True
    paragraph.add_run(" there").underline = True
    fill(document)
    assert [run.text for run in paragraph.runs] == [
        "Hello ",
        "Staff Engineer",
        " there",
    ]
    assert [run.bold for run in paragraph.runs] == [True, None, None]
    assert [run.italic for run in paragraph.runs] == [None, True, None]
    assert [run.underline for run in paragraph.runs] == [None, None, True]


@pytest.mark.parametrize(
    "text",
    [
        "[Company]",
        "[company]",
        "[ROLE]",
        "[Title]",
        "[date]",
        "COMPANY",
        "[ COMPANY ]",
        "[COMPANY",
    ],
)
def test_only_the_exact_tokens_match(text):
    document = docx.Document()
    paragraph = document.add_paragraph(text)
    before = paragraph._p.xml
    assert fill(document) is False
    assert paragraph._p.xml == before


@pytest.mark.parametrize(
    "value",
    [
        "A&B <i>Co</i>",
        r"Back \1 slash \g<0> \\",
        "Dollar $1 ${x} $&",
        "Brackets [TITLE] [DATE] [COMPANY]",
        "<w:t>raw</w:t>",
        "]]> <!-- c --> &amp; &lt;",
        "  padded  ",
        "Zoë Ünïcode 日本語",
    ],
)
def test_a_value_is_inserted_as_literal_text(value):
    document = docx.Document()
    document.add_paragraph("A [COMPANY] B [TITLE] C")
    fill(document, company=value, title=value)
    assert body_texts(reloaded(document)) == [f"A {value} B {value} C"]


def test_a_value_that_looks_like_a_token_is_not_filled_again():
    document = docx.Document()
    document.add_paragraph("[COMPANY] / [TITLE]")
    fill(document, company="[TITLE]", title="[DATE]")
    assert body_texts(document) == ["[TITLE] / [DATE]"]


def test_a_document_with_no_token_is_reported_and_left_as_it_was():
    document = docx.Document()
    document.add_paragraph("Nothing to fill.")
    document.sections[0].header.paragraphs[0].text = "Header"
    before = document.element.xml
    assert fill(document) is False
    assert document.element.xml == before


def test_filling_survives_saving_and_reloading_the_document():
    document = docx.Document()
    paragraph_with_runs(document, "[COM", "PANY] on [DATE]")
    assert body_texts(reloaded(document) if fill(document) else document) == [
        "Acme Widgets on October 10, 2026"
    ]


def test_filling_makes_no_file_system_or_database_call(monkeypatch):
    document = docx.Document()
    document.add_paragraph("[COMPANY] [TITLE] [DATE]")

    def refuse(*_args, **_kwargs):
        raise AssertionError("filling touched the outside world")

    monkeypatch.setattr("builtins.open", refuse)
    monkeypatch.setattr("pathlib.Path.open", refuse)
    monkeypatch.setattr("pathlib.Path.read_bytes", refuse)
    monkeypatch.setattr("pathlib.Path.write_bytes", refuse)
    monkeypatch.setattr(db, "connect", refuse)
    fill(document)
    assert body_texts(document) == ["Acme Widgets Staff Engineer October 10, 2026"]


# --- the packet's copies (jsa packet) -----------------------------------------


@pytest.fixture
def pdb(db_url):
    drop_all_tables(db_url)
    connection = db.connect()
    yield connection
    connection.close()


@pytest.fixture
def env(tmp_path, monkeypatch):
    """A profile whose resume holds the tokens, and its packets directory."""
    profile = copy_example(tmp_path / "profile")
    packets = tmp_path / "packets"
    write_config_toml(
        profile, f'candidate_name = "Pat Example"\npackets_dir = "{packets}"\n'
    )
    monkeypatch.setenv("JSA_PROFILE_DIR", str(profile))
    return profile, packets


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_with_tokens(path, *, header=True):
    document = docx.Document()
    document.add_paragraph("Name line")
    document.add_paragraph("Applying to [COMPANY] as [TITLE] on [DATE].")
    if header:
        document.sections[0].header.paragraphs[0].text = "[COMPANY] header"
    document.save(path)


def packet_folder(packets):
    [folder] = [path for path in packets.iterdir() if path.is_dir()]
    return folder


def copy_in(folder, kind):
    [path] = [p for p in folder.iterdir() if f"_{kind}_" in p.name]
    return path


def english(day: date) -> str:
    return f"{MONTHS[day.month - 1]} {day.day}, {day.year}"


def test_the_resume_copy_is_filled_and_the_profiles_resume_is_byte_for_byte_unchanged(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_tokens(profile / "resume.docx")
    base = digest(profile / "resume.docx")
    seed(pdb, company="Acme Widgets, Inc.", title="Staff Engineer")
    before = datetime.now(ZoneInfo("America/New_York")).date()
    code, _ = jsa_packet(monkeypatch, capsys)
    after = datetime.now(ZoneInfo("America/New_York")).date()
    assert code == 0
    document = docx.Document(str(copy_in(packet_folder(packets), "Resume")))
    assert document.paragraphs[0].text == "Name line"
    assert document.paragraphs[1].text in {
        f"Applying to Acme Widgets as Staff Engineer on {english(day)}."
        for day in (before, after)
    }
    assert document.sections[0].header.paragraphs[0].text == "Acme Widgets header"
    assert digest(profile / "resume.docx") == base


def test_a_docx_cover_letter_copy_is_filled_and_the_profiles_letter_is_unchanged(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_tokens(profile / "cover_letter.docx", header=False)
    base = digest(profile / "cover_letter.docx")
    seed(pdb, title="Data Engineer")
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    letter = docx.Document(str(copy_in(packet_folder(packets), "CoverLetter")))
    assert letter.paragraphs[1].text.startswith(
        "Applying to Acme Widgets as Data Engineer on "
    )
    assert "[" not in letter.paragraphs[1].text
    assert digest(profile / "cover_letter.docx") == base


def test_the_company_filled_is_the_normalized_company_not_the_raw_name(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_tokens(profile / "resume.docx")
    seed(pdb, company="Foo Bar Corp LLC")
    normalized = pdb.execute("SELECT normalized_company FROM postings").fetchone()[0]
    jsa_packet(monkeypatch, capsys)
    document = docx.Document(str(copy_in(packet_folder(packets), "Resume")))
    assert f"Applying to {normalized} as" in document.paragraphs[1].text
    assert "Foo Bar Corp LLC" not in document.paragraphs[1].text


def test_a_source_with_no_token_is_copied_byte_for_byte(pdb, env, monkeypatch, capsys):
    profile, packets = env
    document = docx.Document()
    document.add_paragraph("No placeholders [Company] [ROLE] here.")
    document.save(profile / "resume.docx")
    document.save(profile / "cover_letter.docx")
    seed(pdb)
    jsa_packet(monkeypatch, capsys)
    folder = packet_folder(packets)
    assert (
        copy_in(folder, "Resume").read_bytes() == (profile / "resume.docx").read_bytes()
    )
    assert (
        copy_in(folder, "CoverLetter").read_bytes()
        == (profile / "cover_letter.docx").read_bytes()
    )


@pytest.mark.parametrize("extension", ["pdf", "txt"])
def test_a_cover_letter_that_is_not_a_docx_is_copied_as_it_is(
    pdb, env, monkeypatch, capsys, extension
):
    profile, packets = env
    source = profile / f"cover_letter.{extension}"
    source.write_bytes(b"Dear [COMPANY], about [TITLE] on [DATE]")
    seed(pdb)
    jsa_packet(monkeypatch, capsys)
    assert (
        copy_in(packet_folder(packets), "CoverLetter").read_bytes()
        == source.read_bytes()
    )


def test_the_resume_is_filled_when_the_cover_letter_is_a_pdf(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_tokens(profile / "resume.docx")
    (profile / "cover_letter.pdf").write_bytes(b"[COMPANY]")
    seed(pdb)
    jsa_packet(monkeypatch, capsys)
    folder = packet_folder(packets)
    assert (
        "[COMPANY]"
        not in docx.Document(str(copy_in(folder, "Resume"))).paragraphs[1].text
    )
    assert copy_in(folder, "CoverLetter").read_bytes() == b"[COMPANY]"


def test_values_with_markup_or_regex_characters_land_as_literal_text_in_the_copy(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_tokens(profile / "resume.docx", header=False)
    seed(pdb, company="A&B <i>Co</i> \\1 $1", title="R&D <b>Lead</b> \\g<0> $&")
    normalized = pdb.execute("SELECT normalized_company FROM postings").fetchone()[0]
    jsa_packet(monkeypatch, capsys)
    text = (
        docx.Document(str(copy_in(packet_folder(packets), "Resume"))).paragraphs[1].text
    )
    assert f"Applying to {normalized} as R&D <b>Lead</b> \\g<0> $& on " in text


@pytest.mark.parametrize(
    ("timezone", "other"), [("Pacific/Kiritimati", "Pacific/Pago_Pago")]
)
def test_date_is_today_in_the_profiles_timezone(
    pdb, env, monkeypatch, capsys, timezone, other
):
    profile, packets = env
    save_with_tokens(profile / "resume.docx", header=False)
    write_search_toml(profile, SEARCH_TOML.replace("America/New_York", timezone))
    seed(pdb)
    before = datetime.now(ZoneInfo(timezone)).date()
    jsa_packet(monkeypatch, capsys)
    after = datetime.now(ZoneInfo(timezone)).date()
    text = (
        docx.Document(str(copy_in(packet_folder(packets), "Resume"))).paragraphs[1].text
    )
    assert text.endswith(tuple(f"on {english(day)}." for day in {before, after}))
    # Kiritimati and Pago Pago are 25 hours apart, so the host's own zone can match at most one of them.
    assert not text.endswith(f"on {english(datetime.now(ZoneInfo(other)).date())}.")


def test_date_ignores_the_hosts_tz_variable(pdb, env, monkeypatch, capsys):
    profile, packets = env
    save_with_tokens(profile / "resume.docx", header=False)
    write_search_toml(
        profile, SEARCH_TOML.replace("America/New_York", "Pacific/Kiritimati")
    )
    monkeypatch.setenv("TZ", "Pacific/Pago_Pago")
    seed(pdb)
    before = datetime.now(ZoneInfo("Pacific/Kiritimati")).date()
    jsa_packet(monkeypatch, capsys)
    after = datetime.now(ZoneInfo("Pacific/Kiritimati")).date()
    text = (
        docx.Document(str(copy_in(packet_folder(packets), "Resume"))).paragraphs[1].text
    )
    assert text.endswith(tuple(f"on {english(day)}." for day in {before, after}))


def test_a_packet_for_each_posting_is_filled_with_its_own_company_and_title(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_tokens(profile / "resume.docx", header=False)
    seed(pdb, company="Alpha Labs", title="Backend Engineer")
    seed(pdb, company="Beta Works", title="Data Scientist")
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    texts = {
        docx.Document(str(copy_in(folder, "Resume")))
        .paragraphs[1]
        .text.split(" on ")[0]
        for folder in packets.iterdir()
    }
    assert texts == {
        "Applying to Alpha Labs as Backend Engineer",
        "Applying to Beta Works as Data Scientist",
    }


def test_an_existing_packet_is_not_refilled_when_the_profile_changes(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_tokens(profile / "resume.docx", header=False)
    seed(pdb)
    jsa_packet(monkeypatch, capsys)
    folder = packet_folder(packets)
    copy_path = copy_in(folder, "Resume")
    filled = copy_path.read_bytes()
    save_with_tokens(profile / "resume.docx")
    code, _ = jsa_packet(monkeypatch, capsys)
    assert code == 0
    assert copy_path.read_bytes() == filled


def test_a_users_own_token_typed_into_an_existing_copy_stays_as_typed(
    pdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_tokens(profile / "resume.docx", header=False)
    seed(pdb, jd=JD)
    jsa_packet(monkeypatch, capsys)
    copy_path = copy_in(packet_folder(packets), "Resume")
    revised = docx.Document(str(copy_path))
    revised.add_paragraph("Still to fix: [COMPANY]")
    revised.save(copy_path)
    kept = copy_path.read_bytes()
    jsa_packet(monkeypatch, capsys)
    assert copy_path.read_bytes() == kept


def test_a_dry_run_fills_and_copies_nothing(pdb, env, monkeypatch, capsys):
    profile, packets = env
    save_with_tokens(profile / "resume.docx")
    base = digest(profile / "resume.docx")
    seed(pdb)
    code, _ = jsa_packet(monkeypatch, capsys, "--dry-run")
    assert code == 0
    assert not packets.exists() or not list(packets.iterdir())
    assert digest(profile / "resume.docx") == base
