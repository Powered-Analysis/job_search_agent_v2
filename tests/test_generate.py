"""`jsa generate` (issue #13; PRD 04 "Redline", "Base resume" and "Packet directory"; PRD 02 "Packet queue"; XC-9 to XC-14).

The database is the libSQL container. The outside world is replaced at the
documented seams: HTTP at the transport, Claude at `agent_loop.query`, and `gws` by a stub script
reached through `JSA_GWS_BIN`. The profile and the packets directory are temporary directories.
"""

import asyncio
import json
import re
import shutil
import sys
import threading
import time
import uuid
import zipfile
from pathlib import Path
from types import SimpleNamespace

import docx
import httpx
import pytest
from claude_agent_sdk import ResultMessage
from conftest import drop_all_tables
from docx.oxml import parse_xml
from profile_helpers import (
    ROLE_FAMILIES,
    ROLE_FAMILY,
    copy_example,
    family_dir,
    write_config_toml,
)
from soffice_helpers import rendered
from test_claude_runner import result_message
from test_review import Web
from test_tracker import STUB, appended_rows, calls, params_of

from jsa import agent_loop, cli, db

LONG_AGO = "2020-01-01T00:00:00.000Z"
JD = "# Staff Engineer\n\nBuild the Frobnicator platform.\n"
PLAIN = "Acme Widgets - Staff Engineer"
RESUME_COPY = "PatExample_Resume_StaffEngineer_AcmeWidgets.docx"
RESUME_PDF = RESUME_COPY.replace(".docx", ".pdf")
BASE_LINE = "Built the Quuxlate platform from scratch."
CONFIG = """\
candidate_name = "Pat Example"
tracker_spreadsheet_id = "sheet-123"
packets_dir = "{packets}"

[agents.redline]
model = "claude-opus-5-5"
effort = "medium"
"""
REDLINE_MODEL = "claude-opus-5-5"
REDLINE_EDITS = "redline_edits.json"
# One edit the validator drops (no such paragraph): a well-formed, non-empty redline result.
REDLINE_JSON = json.dumps(
    {
        "edits": [
            {
                "paragraph": 99,
                "find": "nothing here",
                "replace": "nothing there",
                "goal": "ats",
                "jd_quote": "Frobnicator platform",
                "reason": "the same thing",
            }
        ],
        "removals": [],
        "reorders": [],
        "flags": [],
        "explanation": None,
    }
)
NO_EDITS_EXPLANATION = "The resume already uses the posting's terms."

NO_EDITS_RECORD = json.dumps(
    {
        "explanation": NO_EDITS_EXPLANATION,
        "edits": [],
        "removals": [],
        "reorders": [],
        "flags": [],
    }
)


class Agent:
    """Stands in for the SDK's `query`; `respond` maps a redline prompt to the result message the run ends with."""

    def __init__(self):
        self.calls = []
        self.lock = threading.Lock()
        self.in_flight = 0
        self.peak = 0
        self.respond = lambda prompt: result_message(result=REDLINE_JSON)

    async def query(self, *, prompt, options=None, **_ignored):
        with self.lock:
            self.calls.append(SimpleNamespace(prompt=prompt, options=options))
            self.in_flight += 1
            self.peak = max(self.peak, self.in_flight)
        try:
            message = self.respond(prompt)
        finally:
            with self.lock:
                self.in_flight -= 1
        yield message

    @property
    def prompts(self):
        return [call.prompt for call in self.calls]


def failed(**fields) -> ResultMessage:
    return result_message(is_error=True, subtype="error_during_execution", **fields)


@pytest.fixture(autouse=True)
def agent(monkeypatch):
    stand_in = Agent()
    monkeypatch.setattr(agent_loop, "query", stand_in.query)
    return stand_in


@pytest.fixture(autouse=True)
def web(monkeypatch):
    """Every page is live unless a test says otherwise; nothing reaches the network."""
    stand_in = Web()

    def handle(_transport, request):
        return stand_in.handle(request)

    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", handle)
    return stand_in


@pytest.fixture(autouse=True)
def env(tmp_path, monkeypatch):
    """A profile with a distinctive base resume, and its packets directory."""
    profile = copy_example(tmp_path / "profile")
    packets = tmp_path / "packets"
    write_config_toml(profile, CONFIG.format(packets=packets))
    document = docx.Document()
    document.add_paragraph().add_run("Riley Resumeperson").bold = True
    document.add_paragraph(BASE_LINE)
    document.save(family_dir(profile) / "resume.docx")
    monkeypatch.setenv("JSA_PROFILE_DIR", str(profile))
    monkeypatch.delenv("JSA_GENERATE_WORKERS", raising=False)
    return profile, packets


@pytest.fixture(autouse=True)
def gws(tmp_path, monkeypatch):
    """A stub `gws` whose invocations are logged; returns the log's path."""
    stub = tmp_path / "stub-gws"
    stub.write_text(STUB.format(python=sys.executable), encoding="utf-8")
    stub.chmod(0o755)
    log = tmp_path / "gws.log"
    monkeypatch.setenv("JSA_GWS_BIN", str(stub))
    monkeypatch.setenv("STUB_LOG", str(log))
    monkeypatch.delenv("STUB_FAIL_IDS", raising=False)
    monkeypatch.delenv("STUB_FAIL_MODE", raising=False)
    return log


@pytest.fixture
def gdb(db_url):
    """A connection to a database holding no postings but the test's own."""
    drop_all_tables(db_url)
    connection = db.connect()
    yield connection
    connection.close()


def seed(
    conn,
    *,
    company="Acme Widgets, Inc.",
    title="Staff Engineer",
    decision="Apply",
    tracked=False,
    closed=False,
    jd=JD,
    role_family=ROLE_FAMILY,
):
    posting_id = db.insert_posting(
        conn,
        company=company,
        title=title,
        url=f"https://careers.example.com/jobs/{uuid.uuid4().hex}",
        search_agent="claude",
    )
    conn.execute(
        "UPDATE postings SET decision = ?, decided_at = ?, added_to_tracker = ?, "
        "closed_at = ?, jd_markdown = ?, role_family = ? WHERE id = ?",
        (
            decision,
            LONG_AGO if decision else None,
            int(tracked),
            LONG_AGO if closed else None,
            jd,
            role_family if decision == "Apply" else None,
            posting_id,
        ),
    )
    return posting_id


def url_of(conn, posting_id) -> str:
    return conn.execute(
        "SELECT url FROM postings WHERE id = ?", (posting_id,)
    ).fetchone()[0]


def column(conn, posting_id, name):
    return conn.execute(
        f"SELECT {name} FROM postings WHERE id = ?", (posting_id,)
    ).fetchone()[0]


def jsa(monkeypatch, capsys, command, *args):
    monkeypatch.setattr(sys, "argv", ["jsa", command, *map(str, args)])
    code = 0
    try:
        cli.main()
    except SystemExit as exit_:
        code = exit_.code if isinstance(exit_.code, int) else 1
    captured = capsys.readouterr()
    return code, captured.out + captured.err


def jsa_generate(monkeypatch, capsys, *args):
    return jsa(monkeypatch, capsys, "generate", *args)


def entries(directory: Path) -> set[str]:
    return {path.name for path in directory.iterdir()} if directory.exists() else set()


def edit_copy(path: Path, text: str) -> None:
    document = docx.Document(str(path))
    document.add_paragraph(text)
    document.save(path)


def numbered(conn, count, **fields):
    return [seed(conn, company=f"Org{n} Labs", **fields) for n in range(count)]


# --- the happy path -----------------------------------------------------------


def test_an_open_apply_posting_gets_a_full_packet_and_a_tracker_row(
    gdb, env, gws, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    folder = packets / PLAIN
    assert entries(folder) == {
        "job_posting.md",
        RESUME_COPY,
        RESUME_PDF,
        REDLINE_EDITS,
    }
    assert (folder / "job_posting.md").read_text(encoding="utf-8") == JD
    assert (folder / RESUME_COPY).read_bytes() == (
        family_dir(profile) / "resume.docx"
    ).read_bytes()
    assert [row[0] for row in appended_rows(gws)] == [posting_id]
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_a_second_run_finds_nothing_to_do(gdb, agent, gws, monkeypatch, capsys):
    seed(gdb)
    jsa_generate(monkeypatch, capsys)
    agent.calls.clear()
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert agent.calls == []
    assert len(calls(gws)) == 1


def test_the_tracker_row_is_what_jsa_track_would_have_written(
    gdb, gws, tmp_path, monkeypatch, capsys
):
    first = seed(gdb, company="Beta Labs")
    second = seed(gdb, company="Gamma Labs")
    jsa_generate(monkeypatch, capsys, "--id", first)
    generated = appended_rows(gws)
    gws.unlink()
    jsa(monkeypatch, capsys, "track", "--id", second)
    tracked = appended_rows(gws)
    assert len(generated) == len(tracked) == 1
    # Same A:H shape: the ID, company, title, URL, then the blank user-owned cells.
    assert len(generated[0]) == len(tracked[0])
    assert generated[0][2] == tracked[0][2] == "Staff Engineer"
    assert generated[0][6:] == tracked[0][6:]


@pytest.mark.parametrize(
    "state",
    [{"decision": "Skip"}, {"decision": None}, {"tracked": True}, {"closed": True}],
    ids=["skip", "undecided", "tracked", "closed"],
)
def test_only_apply_untracked_open_postings_are_queued(
    gdb, env, agent, state, monkeypatch, capsys
):
    _, packets = env
    seed(gdb, **state)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert entries(packets) == set()
    assert agent.calls == []


def test_an_id_for_a_non_apply_posting_builds_nothing(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb, decision="Skip")
    jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert entries(packets) == set()
    assert agent.calls == [] and not gws.exists()


def test_each_row_is_built_and_tracked_with_its_own_redline(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    ids = [
        seed(gdb, company="Beta Labs", jd="# Role\n\nBETA-JD-MARKER\n"),
        seed(gdb, company="Gamma Labs", jd="# Role\n\nGAMMA-JD-MARKER\n"),
    ]
    agent.respond = lambda prompt: redline_reply(
        explanation="for Beta" if "BETA-JD-MARKER" in prompt else "for Gamma"
    )(prompt)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    beta = packets / "Beta Labs - Staff Engineer"
    gamma = packets / "Gamma Labs - Staff Engineer"
    assert proposed(beta)["explanation"] == "for Beta"
    assert proposed(gamma)["explanation"] == "for Gamma"
    assert all(column(gdb, i, "added_to_tracker") == 1 for i in ids)
    assert len(calls(gws)) == 2


def test_two_postings_sharing_a_name_get_separate_packets(
    gdb, env, monkeypatch, capsys
):
    _, packets = env
    seed(gdb)
    second = seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert entries(packets) == {PLAIN, f"{PLAIN} ({second})"}
    for name in entries(packets):
        assert REDLINE_EDITS in entries(packets / name)


# --- the liveness re-check ----------------------------------------------------


def test_a_posting_found_closed_is_marked_closed_and_nothing_is_built(
    gdb, env, web, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    web.gone(url_of(gdb, posting_id))
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert column(gdb, posting_id, "closed_at") is not None
    assert column(gdb, posting_id, "added_to_tracker") == 0
    assert entries(packets) == set()
    assert agent.calls == []
    assert not gws.exists()


def test_the_report_counts_the_postings_that_closed(gdb, env, web, monkeypatch, capsys):
    _, packets = env
    dead = [seed(gdb, company="Dead One"), seed(gdb, company="Dead Two")]
    seed(gdb, company="Alive Labs")
    for posting_id in dead:
        web.gone(url_of(gdb, posting_id))
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert re.search(r"\b2\b[^\n]*clos", output), output
    assert entries(packets) == {"Alive Labs - Staff Engineer"}


def test_the_closed_count_is_reported_even_when_it_is_zero(gdb, monkeypatch, capsys):
    seed(gdb)
    _, output = jsa_generate(monkeypatch, capsys)
    assert re.search(r"\b0\b[^\n]*clos", output), output


def test_an_unreachable_posting_is_not_closed_and_is_still_built(
    gdb, env, web, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    web.blocked(url_of(gdb, posting_id))
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert column(gdb, posting_id, "closed_at") is None
    assert REDLINE_EDITS in entries(packets / PLAIN)


def test_an_id_builds_a_closed_posting_anyway(gdb, env, web, gws, monkeypatch, capsys):
    _, packets = env
    posting_id = seed(gdb, closed=True)
    web.gone(url_of(gdb, posting_id))
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert REDLINE_EDITS in entries(packets / PLAIN)
    assert len(calls(gws)) == 1
    assert column(gdb, posting_id, "added_to_tracker") == 1


# --- re-entering a packet -----------------------------------------------------


def test_a_bare_directory_left_by_jsa_packet_is_completed_not_skipped(
    gdb, env, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    jsa(monkeypatch, capsys, "packet")
    assert REDLINE_EDITS not in entries(packets / PLAIN)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert entries(packets / PLAIN) == {
        "job_posting.md",
        RESUME_COPY,
        RESUME_PDF,
        REDLINE_EDITS,
    }
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_an_empty_directory_is_completed_with_the_job_posting_and_resume_copy(
    gdb, env, monkeypatch, capsys
):
    _, packets = env
    seed(gdb)
    (packets / PLAIN).mkdir(parents=True)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert entries(packets / PLAIN) == {
        "job_posting.md",
        RESUME_COPY,
        RESUME_PDF,
        REDLINE_EDITS,
    }


def test_an_id_on_a_tracked_posting_builds_the_redline_but_adds_no_second_row(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb, tracked=True)
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert REDLINE_EDITS in entries(packets / PLAIN)
    assert not gws.exists()


# --- job_posting.md -----------------------------------------------------------


def test_the_queue_path_never_overwrites_an_existing_job_posting(
    gdb, env, monkeypatch, capsys
):
    _, packets = env
    seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / "job_posting.md").write_text("hand edited", encoding="utf-8")
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert (folder / "job_posting.md").read_text(encoding="utf-8") == "hand edited"


def test_an_id_overwrites_the_job_posting_with_the_stored_description(
    gdb, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / "job_posting.md").write_text("stale text", encoding="utf-8")
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert (folder / "job_posting.md").read_text(encoding="utf-8") == JD
    assert "Frobnicator" in agent.prompts[0]
    assert "stale text" not in agent.prompts[0]


def test_an_id_on_a_null_description_never_overwrites_a_hand_filled_file(
    gdb, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb, jd=None)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / "job_posting.md").write_text(
        "Zebrafish wrangler duties", encoding="utf-8"
    )
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert (folder / "job_posting.md").read_text(encoding="utf-8") == (
        "Zebrafish wrangler duties"
    )
    assert "Zebrafish wrangler duties" in agent.prompts[0]


# --- never review blind -------------------------------------------------------


def test_a_null_jd_without_a_job_posting_gets_its_head_only_and_stays_queued(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb, jd=None)
    _, output = jsa_generate(monkeypatch, capsys)
    assert entries(packets / PLAIN) == {RESUME_COPY}
    assert agent.calls == []
    assert not gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 0
    assert str(posting_id) in output


def test_a_null_jd_flagged_on_its_own_leaves_the_exit_status_zero(
    gdb, monkeypatch, capsys
):
    seed(gdb, jd=None)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0


def test_a_hand_filled_job_posting_is_the_jd_and_is_left_unchanged(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb, jd=None)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / "job_posting.md").write_text(
        "Zebrafish wrangler duties", encoding="utf-8"
    )
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert "Zebrafish wrangler duties" in agent.prompts[0]
    assert (folder / "job_posting.md").read_text(encoding="utf-8") == (
        "Zebrafish wrangler duties"
    )
    assert REDLINE_EDITS in entries(folder)
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_a_blind_row_does_not_stop_the_others_and_is_not_a_failure(
    gdb, env, gws, monkeypatch, capsys
):
    _, packets = env
    blind = seed(gdb, company="Blind Labs", jd=None)
    sighted = seed(gdb, company="Sighted Labs")
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert REDLINE_EDITS not in entries(packets / "Blind Labs - Staff Engineer")
    assert column(gdb, blind, "added_to_tracker") == 0
    assert column(gdb, sighted, "added_to_tracker") == 1


# --- what the agent is given --------------------------------------------------


def test_the_prompt_carries_no_other_profile_content(
    gdb, env, agent, monkeypatch, capsys
):
    profile, _ = env
    seed(gdb)
    jsa_generate(monkeypatch, capsys)
    (prompt,) = agent.prompts
    for path in (profile / "search").glob("*.md"):
        for line in path.read_text(encoding="utf-8").splitlines():
            if len(line.strip()) > 30:
                assert line.strip() not in prompt, path.name
    for private in ("Pat Example", "Jordan Example", "sheet-123"):
        assert private not in prompt


def test_the_base_resume_changing_later_does_not_reach_an_existing_copys_redline(
    gdb, env, agent, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb)
    jsa_generate(monkeypatch, capsys)
    (packets / PLAIN / REDLINE_EDITS).unlink()
    document = docx.Document()
    document.add_paragraph("BASE ONLY: a different base resume entirely")
    document.save(family_dir(profile) / "resume.docx")
    agent.calls.clear()
    jsa_generate(monkeypatch, capsys, "--id", posting_id)
    (prompt,) = agent.prompts
    assert "BASE ONLY" not in prompt
    assert BASE_LINE in prompt


# --- the resume copy ----------------------------------------------------------


def test_an_existing_resume_copy_is_never_modified(gdb, env, monkeypatch, capsys):
    _, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / RESUME_COPY).write_bytes(b"not even a docx, just my revision")
    jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert (folder / RESUME_COPY).read_bytes() == b"not even a docx, just my revision"


def test_a_deleted_resume_copy_is_restored_from_the_base_before_the_redline(
    gdb, env, agent, monkeypatch, capsys
):
    profile, packets = env
    seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / "job_posting.md").write_text(JD, encoding="utf-8")
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert (folder / RESUME_COPY).read_bytes() == (
        family_dir(profile) / "resume.docx"
    ).read_bytes()
    assert BASE_LINE in agent.prompts[0]


# --- the cover letter (issue #112; PRD 04 "Cover letter") ---------------------

COVER_LETTER_COPY = RESUME_COPY.replace("_Resume_", "_CoverLetter_")
COVER_LETTER_PDF = COVER_LETTER_COPY.replace(".docx", ".pdf")
COVER_LETTER_TEXT = b"Dear Hiring Manager, ZEBRAQUILL-MARKER"


def test_generate_puts_the_cover_letter_copy_in_the_packet_head(
    gdb, env, monkeypatch, capsys
):
    profile, packets = env
    (family_dir(profile) / "cover_letter.docx").write_bytes(COVER_LETTER_TEXT)
    seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert entries(packets / PLAIN) == {
        "job_posting.md",
        RESUME_COPY,
        RESUME_PDF,
        COVER_LETTER_COPY,
        COVER_LETTER_PDF,
        REDLINE_EDITS,
    }
    assert (packets / PLAIN / COVER_LETTER_COPY).read_bytes() == COVER_LETTER_TEXT


def test_a_docx_cover_letter_copy_gets_its_pdf_only_when_it_has_none(
    gdb, env, soffice, monkeypatch, capsys
):
    profile, packets = env
    (family_dir(profile) / "cover_letter.docx").write_bytes(COVER_LETTER_TEXT)
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / COVER_LETTER_PDF).write_bytes(b"MY OWN EXPORT")
    jsa_generate(monkeypatch, capsys)
    assert (folder / COVER_LETTER_PDF).read_bytes() == b"MY OWN EXPORT"
    (folder / COVER_LETTER_PDF).unlink()
    jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert (folder / COVER_LETTER_PDF).is_file()
    assert COVER_LETTER_COPY in rendered(soffice)


def test_generate_without_a_cover_letter_makes_none(gdb, env, monkeypatch, capsys):
    _, packets = env
    seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert not [name for name in entries(packets / PLAIN) if "CoverLetter" in name]


def test_an_existing_cover_letter_copy_is_never_modified(gdb, env, monkeypatch, capsys):
    profile, packets = env
    (family_dir(profile) / "cover_letter.docx").write_bytes(COVER_LETTER_TEXT)
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / COVER_LETTER_COPY).write_bytes(b"my edited cover letter")
    for args in ((), ("--id", posting_id)):
        jsa_generate(monkeypatch, capsys, *args)
        assert (folder / COVER_LETTER_COPY).read_bytes() == b"my edited cover letter"


def test_a_deleted_cover_letter_copy_is_restored_from_the_profile(
    gdb, env, monkeypatch, capsys
):
    profile, packets = env
    (family_dir(profile) / "cover_letter.pdf").write_bytes(COVER_LETTER_TEXT)
    seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    copy = folder / COVER_LETTER_COPY.replace(".docx", ".pdf")
    assert copy.read_bytes() == COVER_LETTER_TEXT


def test_the_redline_does_not_see_the_cover_letter(
    gdb, env, agent, monkeypatch, capsys
):
    profile, _ = env
    document = docx.Document()
    document.add_paragraph("ZEBRAQUILL-MARKER")
    document.save(family_dir(profile) / "cover_letter.docx")
    seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert agent.calls
    assert not [prompt for prompt in agent.prompts if "ZEBRAQUILL" in prompt]


def test_more_than_one_cover_letter_fails_before_any_row_is_processed(
    gdb, env, agent, gws, monkeypatch, capsys
):
    profile, packets = env
    (family_dir(profile) / "cover_letter.docx").write_bytes(COVER_LETTER_TEXT)
    (family_dir(profile) / "cover_letter.pdf").write_bytes(COVER_LETTER_TEXT)
    seed(gdb)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert "cover_letter.docx" in output and "cover_letter.pdf" in output
    assert not packets.exists() or not list(packets.iterdir())
    assert agent.calls == []
    assert not gws.exists()


# --- failures -----------------------------------------------------------------


def test_an_api_error_result_flags_the_row_with_its_status_and_exits_non_zero(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    agent.respond = lambda prompt: failed(api_error_status=529, result="overloaded")
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert "529" in output
    assert REDLINE_EDITS not in entries(packets / PLAIN)
    assert not gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 0


@pytest.mark.parametrize("text", ["", "   \n"], ids=["empty", "blank"])
def test_an_empty_result_flags_the_row_and_exits_non_zero(
    gdb, env, agent, gws, text, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    agent.respond = lambda prompt: result_message(result=text)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert not (packets / PLAIN / REDLINE_EDITS).exists()
    assert not gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 0


def test_one_failed_row_does_not_stop_the_rest_but_the_exit_is_non_zero(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    bad = seed(gdb, company="Bad Labs", jd="# Role\n\nBAD-JD-MARKER\n")
    good = seed(gdb, company="Good Labs")
    agent.respond = lambda prompt: (
        failed(api_error_status=500)
        if "BAD-JD-MARKER" in prompt
        else result_message(result=REDLINE_JSON)
    )
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(bad) in output
    assert column(gdb, good, "added_to_tracker") == 1
    assert column(gdb, bad, "added_to_tracker") == 0
    assert REDLINE_EDITS in entries(packets / "Good Labs - Staff Engineer")


@pytest.mark.parametrize("mode", ["exit", "garbage", "no-rows"])
def test_a_failed_track_exits_non_zero_leaves_the_row_untracked_and_rolls_nothing_back(
    gdb, env, gws, mode, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    monkeypatch.setenv("STUB_FAIL_IDS", str(posting_id))
    monkeypatch.setenv("STUB_FAIL_MODE", mode)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(posting_id) in output
    assert column(gdb, posting_id, "added_to_tracker") == 0
    assert entries(packets / PLAIN) == {
        "job_posting.md",
        RESUME_COPY,
        RESUME_PDF,
        REDLINE_EDITS,
    }


def test_a_failed_track_leaves_the_redline_so_a_rerun_resumes_at_the_append(
    gdb, env, agent, gws, monkeypatch, capsys
):
    posting_id = seed(gdb)
    monkeypatch.setenv("STUB_FAIL_IDS", str(posting_id))
    jsa_generate(monkeypatch, capsys)
    monkeypatch.delenv("STUB_FAIL_IDS")
    agent.calls.clear()
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert agent.calls == []
    assert column(gdb, posting_id, "added_to_tracker") == 1


# --- concurrency --------------------------------------------------------------


def test_three_redline_runs_are_in_flight_together_by_default(
    gdb, agent, monkeypatch, capsys
):
    numbered(gdb, 3)
    barrier = threading.Barrier(3, timeout=20)

    def respond(prompt):
        try:
            barrier.wait()
        except threading.BrokenBarrierError:
            return failed(result="the other runs never started")
        return result_message(result=REDLINE_JSON)

    agent.respond = respond
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert agent.peak == 3


@pytest.mark.parametrize("workers", [None, "1", "2"], ids=["default", "one", "two"])
def test_no_more_than_the_configured_number_of_runs_are_in_flight(
    gdb, agent, workers, monkeypatch, capsys
):
    if workers:
        monkeypatch.setenv("JSA_GENERATE_WORKERS", workers)
    limit = int(workers or 3)
    ids = numbered(gdb, 7)

    def respond(prompt):
        time.sleep(0.1)
        return result_message(result=REDLINE_JSON)

    agent.respond = respond
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert len(agent.calls) == 7
    assert agent.peak <= limit
    assert all(column(gdb, i, "added_to_tracker") == 1 for i in ids)


def test_tracker_appends_never_overlap(gdb, agent, tmp_path, monkeypatch, capsys):
    numbered(gdb, 3)
    log = tmp_path / "overlap.log"
    stub = tmp_path / "slow-gws"
    stub.write_text(
        f"#!{sys.executable}\n"
        "import json, sys, time\n"
        f"log = open({str(log)!r}, 'a')\n"
        "log.write('start\\n'); log.flush()\n"
        "time.sleep(0.3)\n"
        "log.write('end\\n'); log.flush()\n"
        "print(json.dumps({'updates': {'updatedRows': 1}}))\n",
        encoding="utf-8",
    )
    stub.chmod(0o755)
    monkeypatch.setenv("JSA_GWS_BIN", str(stub))
    barrier = threading.Barrier(3, timeout=20)

    def respond(prompt):
        try:
            barrier.wait()
        except threading.BrokenBarrierError:
            pass
        return result_message(result=REDLINE_JSON)

    agent.respond = respond
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert log.read_text().split() == ["start", "end"] * 3


def test_each_finished_row_is_appended_right_away_not_at_the_end(
    gdb, agent, gws, monkeypatch, capsys
):
    fast = seed(gdb, company="FastLabs")
    slow = seed(gdb, company="SlowLabs", jd="# Role\n\nSLOW-JD-MARKER\n")

    def respond(prompt):
        if "SLOW-JD-MARKER" in prompt:
            deadline = time.monotonic() + 30
            while not gws.exists() and time.monotonic() < deadline:
                time.sleep(0.05)
            if not gws.exists():
                return failed(result="the fast row was never appended")
        return result_message(result=REDLINE_JSON)

    agent.respond = respond
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert column(gdb, fast, "added_to_tracker") == 1
    assert column(gdb, slow, "added_to_tracker") == 1
    assert appended_rows(gws)[0][0] == fast


# --- --dry-run ----------------------------------------------------------------


def test_a_dry_run_names_the_folders_and_does_nothing_else(
    gdb, env, web, agent, gws, monkeypatch, capsys
):
    _, packets = env
    open_id = seed(gdb, company="Open Labs")
    dead_id = seed(gdb, company="Dead Labs")
    web.gone(url_of(gdb, dead_id))
    code, output = jsa_generate(monkeypatch, capsys, "--dry-run")
    assert code == 0
    assert "Open Labs - Staff Engineer" in output
    assert agent.calls == []
    assert not gws.exists()
    assert not packets.exists() or entries(packets) == set()
    for posting_id in (open_id, dead_id):
        assert column(gdb, posting_id, "added_to_tracker") == 0
        assert column(gdb, posting_id, "closed_at") is None


def test_a_dry_run_with_an_id_writes_nothing_even_to_an_existing_packet(
    gdb, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / "job_posting.md").write_text("stale text", encoding="utf-8")
    code, _ = jsa_generate(monkeypatch, capsys, "--dry-run", "--id", posting_id)
    assert code == 0
    assert entries(folder) == {"job_posting.md"}
    assert (folder / "job_posting.md").read_text(encoding="utf-8") == "stale text"
    assert agent.calls == []


# --- requiredness -------------------------------------------------------------


def _break_resume(profile, state):
    if state == "missing":
        (family_dir(profile) / "resume.docx").unlink()
    else:
        (family_dir(profile) / "resume.docx").write_bytes(b"")


def _break_tracker_id(profile, state):
    packets = profile.parent / "packets"
    text = CONFIG.format(packets=packets)
    if state == "missing":
        text = text.replace('tracker_spreadsheet_id = "sheet-123"\n', "")
    else:
        text = text.replace("sheet-123", "")
    write_config_toml(profile, text)


def _break_redline_table(profile, state):
    packets = profile.parent / "packets"
    write_config_toml(
        profile, CONFIG.format(packets=packets).split("[agents.redline]")[0]
    )


@pytest.mark.parametrize(
    "breaker, state",
    [
        (_break_resume, "missing"),
        (_break_resume, "empty"),
        (_break_tracker_id, "missing"),
        (_break_tracker_id, "empty"),
        (_break_redline_table, "missing"),
    ],
    ids=[
        "resume-missing",
        "resume-empty",
        "tracker-id-missing",
        "tracker-id-empty",
        "redline-agent-missing",
    ],
)
def test_a_missing_requirement_fails_before_any_row_is_processed(
    gdb, env, web, agent, gws, breaker, state, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb)
    breaker(profile, state)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert "profile.example" in output
    assert entries(packets) == set()
    assert agent.calls == []
    assert not gws.exists()
    assert web.requests == []
    assert column(gdb, posting_id, "closed_at") is None
    assert column(gdb, posting_id, "added_to_tracker") == 0


def test_the_requirements_apply_even_when_the_queue_is_empty(
    gdb, env, monkeypatch, capsys
):
    profile, _ = env
    (family_dir(profile) / "resume.docx").unlink()
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert "profile.example" in output


@pytest.mark.parametrize("family", ROLE_FAMILIES)
@pytest.mark.parametrize("state", ["missing", "empty"])
def test_a_bad_resume_in_any_role_family_fails_before_any_row_is_processed_naming_its_folder(
    gdb, env, agent, gws, family, state, monkeypatch, capsys
):
    profile, packets = env
    seed(gdb)
    resume = family_dir(profile, family) / "resume.docx"
    if state == "missing":
        resume.unlink()
    else:
        resume.write_bytes(b"")
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(family_dir(profile, family)) in output
    assert "profile.example" in output
    assert entries(packets) == set()
    assert agent.calls == [] and not gws.exists()


def test_a_profile_without_role_families_fails_before_any_row_is_processed(
    gdb, env, agent, gws, monkeypatch, capsys
):
    profile, packets = env
    seed(gdb)
    shutil.rmtree(profile / "role_families")
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert "profile.example" in output
    assert entries(packets) == set()
    assert agent.calls == [] and not gws.exists()


def test_two_cover_letters_in_any_role_family_fail_the_run_naming_its_folder(
    gdb, env, agent, gws, monkeypatch, capsys
):
    profile, packets = env
    other = family_dir(profile, ROLE_FAMILIES[1])
    seed(gdb)
    (other / "cover_letter.docx").write_bytes(COVER_LETTER_TEXT)
    (other / "cover_letter.pdf").write_bytes(COVER_LETTER_TEXT)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(other) in output
    assert entries(packets) == set()
    assert agent.calls == [] and not gws.exists()


# --- role families (PRD 04 "Role families") -----------------------------------


def tabs_appended_to(gws) -> list[str]:
    return [params_of(argv)["range"].split("!")[0] for argv in calls(gws)]


@pytest.mark.parametrize(
    "family", [None, "Retired Role"], ids=["none", "not-in-the-profile"]
)
def test_a_row_without_a_usable_role_family_is_flagged_and_nothing_is_built_or_tracked(
    gdb, env, agent, gws, family, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb, role_family=family)
    code, output = jsa_generate(monkeypatch, capsys)
    assert str(posting_id) in output
    assert f"jsa revise --id {posting_id}" in output
    assert entries(packets) == set()
    assert agent.calls == [] and not gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 0
    again, _ = jsa_generate(monkeypatch, capsys)
    assert agent.calls == [] and not gws.exists()
    assert code == again
    assert [
        row[0]
        for row in gdb.execute(
            "SELECT id FROM postings WHERE decision = 'Apply' AND added_to_tracker = 0"
        ).fetchall()
    ] == [posting_id]


def test_a_row_without_a_role_family_does_not_stop_the_others(
    gdb, env, gws, monkeypatch, capsys
):
    _, packets = env
    orphan = seed(gdb, company="Orphan Labs", role_family=None)
    served = seed(gdb, company="Served Labs")
    jsa_generate(monkeypatch, capsys)
    assert entries(packets) == {"Served Labs - Staff Engineer"}
    assert [row[0] for row in appended_rows(gws)] == [served]
    assert column(gdb, orphan, "added_to_tracker") == 0
    assert column(gdb, served, "added_to_tracker") == 1


def test_an_id_for_a_row_without_a_role_family_builds_nothing(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb, role_family=None)
    _, output = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert f"jsa revise --id {posting_id}" in output
    assert entries(packets) == set()
    assert agent.calls == [] and not gws.exists()


def test_each_packet_starts_from_its_own_role_familys_files_and_tracks_on_its_tab(
    gdb, env, agent, gws, monkeypatch, capsys
):
    profile, packets = env
    first, second = ROLE_FAMILIES
    document = docx.Document()
    document.add_paragraph("SECOND FAMILY BASE LINE")
    document.save(family_dir(profile, second) / "resume.docx")
    (family_dir(profile, second) / "cover_letter.docx").write_bytes(COVER_LETTER_TEXT)
    a = seed(gdb, company="Alpha Labs", role_family=first)
    b = seed(gdb, company="Beta Labs", role_family=second)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    alpha, beta = (
        packets / "Alpha Labs - Staff Engineer",
        packets / "Beta Labs - Staff Engineer",
    )
    assert (alpha / "PatExample_Resume_StaffEngineer_AlphaLabs.docx").read_bytes() == (
        family_dir(profile, first) / "resume.docx"
    ).read_bytes()
    assert (beta / "PatExample_Resume_StaffEngineer_BetaLabs.docx").read_bytes() == (
        family_dir(profile, second) / "resume.docx"
    ).read_bytes()
    assert not any("CoverLetter" in name for name in entries(alpha))
    assert any("CoverLetter" in name for name in entries(beta))
    # The redline reads each posting's own packet copy, so each family's base line reaches exactly one prompt.
    assert sum(BASE_LINE in prompt for prompt in agent.prompts) == 1
    assert any(
        "SECOND FAMILY BASE LINE" in prompt and BASE_LINE not in prompt
        for prompt in agent.prompts
    )
    ids_to_tabs = dict(
        zip([row[0] for row in appended_rows(gws)], tabs_appended_to(gws))
    )
    assert ids_to_tabs == {a: f"'{first}'", b: f"'{second}'"}


def test_a_missing_tab_fails_that_row_naming_the_tab_and_the_others_continue(
    gdb, env, gws, monkeypatch, capsys
):
    absent = ROLE_FAMILIES[1]
    monkeypatch.setenv("STUB_MISSING_TABS", absent)
    stranded = seed(gdb, company="Stranded Labs", role_family=absent)
    served = seed(gdb, company="Served Labs")
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert f"failed: posting {stranded}:" in output
    assert absent in output
    assert column(gdb, stranded, "added_to_tracker") == 0
    assert column(gdb, served, "added_to_tracker") == 1


def test_a_row_is_marked_tracked_only_after_a_confirmed_append(
    gdb, env, gws, monkeypatch, capsys
):
    posting_id = seed(gdb)
    monkeypatch.setenv("STUB_FAIL_IDS", str(posting_id))
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert column(gdb, posting_id, "added_to_tracker") == 0
    monkeypatch.delenv("STUB_FAIL_IDS")
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_a_resume_and_cover_letter_at_the_top_of_the_profile_are_never_read(
    gdb, env, gws, monkeypatch, capsys
):
    profile, packets = env
    (profile / "resume.docx").write_bytes(b"TOP LEVEL RESUME")
    (profile / "cover_letter.docx").write_bytes(b"TOP LEVEL LETTER")
    seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert (packets / PLAIN / RESUME_COPY).read_bytes() == (
        family_dir(profile) / "resume.docx"
    ).read_bytes()
    assert not [name for name in entries(packets / PLAIN) if "CoverLetter" in name]


# --- a changed role family never adds files to an existing packet (PRD 04 "Edge cases") ---

OLD_FAMILY, NEW_FAMILY = ROLE_FAMILIES
LETTER_COPY_PDF = COVER_LETTER_COPY.replace(".docx", ".pdf")


def packet_of_the_old_family(profile, packets, *, letter=None):
    """A packet as the old role family left it, with the user's own edit in the resume copy."""
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / "job_posting.md").write_text(JD, encoding="utf-8")
    shutil.copy(family_dir(profile, OLD_FAMILY) / "resume.docx", folder / RESUME_COPY)
    edit_copy(folder / RESUME_COPY, "MY OWN EDIT")
    if letter:
        (folder / letter).write_bytes(b"THE OLD FAMILYS LETTER")
    return folder


@pytest.mark.parametrize(
    "old_letter, new_letter",
    [
        (None, "cover_letter.docx"),
        (LETTER_COPY_PDF, "cover_letter.docx"),
        (None, "cover_letter.pdf"),
    ],
    ids=["old-had-none", "old-was-a-pdf-new-is-a-docx", "old-had-none-new-is-a-pdf"],
)
@pytest.mark.parametrize("by_id", [False, True], ids=["queue", "id"])
def test_a_changed_role_family_adds_neither_its_resume_nor_its_cover_letter_to_an_existing_packet(
    gdb, env, by_id, old_letter, new_letter, monkeypatch, capsys
):
    profile, packets = env
    (family_dir(profile, NEW_FAMILY) / new_letter).write_bytes(b"NEW FAMILY LETTER")
    document = docx.Document()
    document.add_paragraph("NEW FAMILY BASE LINE")
    document.save(family_dir(profile, NEW_FAMILY) / "resume.docx")
    posting_id = seed(gdb, role_family=NEW_FAMILY)
    folder = packet_of_the_old_family(profile, packets, letter=old_letter)
    before = {name: (folder / name).read_bytes() for name in entries(folder)}
    code, _ = jsa_generate(
        monkeypatch, capsys, *(["--id", posting_id] if by_id else [])
    )
    assert code == 0
    after = entries(folder)
    assert after - set(before) <= {RESUME_PDF, REDLINE_EDITS}
    assert {name for name in after if "CoverLetter" in name} == (
        {old_letter} if old_letter else set()
    )
    assert (folder / RESUME_COPY).read_bytes() == before[RESUME_COPY]
    for name in after:
        assert b"NEW FAMILY LETTER" not in (folder / name).read_bytes()
        assert b"NEW FAMILY BASE LINE" not in (folder / name).read_bytes()
    if old_letter:
        assert (folder / old_letter).read_bytes() == before[old_letter]


# --- a deleted cover letter PDF is re-made whichever role family the packet's letter came from (issue #214; PRD 04 "PDF copies") ---

LETTER_COPY_DOCX = COVER_LETTER_COPY


def packet_with_letter(profile, packets, letter, *, pdf=True):
    """An old role family's packet whose cover letter copy is `letter`; its PDF is there unless `pdf` is false."""
    folder = packet_of_the_old_family(profile, packets, letter=letter)
    if pdf:
        (folder / LETTER_COPY_PDF).write_bytes(b"MY OWN EXPORT")
    return folder


@pytest.mark.parametrize(
    "new_letter",
    [None, "cover_letter.pdf", "cover_letter.txt", "cover_letter.docx"],
    ids=["none", "pdf", "txt", "docx"],
)
@pytest.mark.parametrize("by_id", [False, True], ids=["queue", "id"])
def test_a_docx_cover_letter_in_the_packet_gets_its_deleted_pdf_whatever_the_current_role_familys_letter(
    gdb, env, soffice, by_id, new_letter, monkeypatch, capsys
):
    profile, packets = env
    if new_letter:
        (family_dir(profile, NEW_FAMILY) / new_letter).write_bytes(b"NEW FAMILY LETTER")
    posting_id = seed(gdb, role_family=NEW_FAMILY)
    folder = packet_with_letter(profile, packets, LETTER_COPY_DOCX, pdf=False)
    code, _ = jsa_generate(
        monkeypatch, capsys, *(["--id", posting_id] if by_id else [])
    )
    assert code == 0
    assert (folder / LETTER_COPY_PDF).is_file()
    assert (folder / LETTER_COPY_DOCX).read_bytes() == b"THE OLD FAMILYS LETTER"
    assert rendered(soffice).count(LETTER_COPY_DOCX) == 1


@pytest.mark.parametrize(
    "new_letter", [None, "cover_letter.docx"], ids=["none", "docx"]
)
def test_an_existing_pdf_beside_an_earlier_role_familys_docx_cover_letter_is_never_replaced(
    gdb, env, soffice, new_letter, monkeypatch, capsys
):
    profile, packets = env
    if new_letter:
        (family_dir(profile, NEW_FAMILY) / new_letter).write_bytes(b"NEW FAMILY LETTER")
    posting_id = seed(gdb, role_family=NEW_FAMILY)
    folder = packet_with_letter(profile, packets, LETTER_COPY_DOCX)
    for args in ((), ("--id", posting_id)):
        code, _ = jsa_generate(monkeypatch, capsys, *args)
        assert code == 0
        assert (folder / LETTER_COPY_PDF).read_bytes() == b"MY OWN EXPORT"
    assert COVER_LETTER_COPY not in rendered(soffice)


@pytest.mark.parametrize("letter", ["pdf", "txt"])
def test_a_cover_letter_in_another_format_in_the_packet_is_never_rendered(
    gdb, env, soffice, letter, monkeypatch, capsys
):
    profile, packets = env
    (family_dir(profile, NEW_FAMILY) / "cover_letter.docx").write_bytes(b"NEW LETTER")
    posting_id = seed(gdb, role_family=NEW_FAMILY)
    name = COVER_LETTER_COPY.replace(".docx", f".{letter}")
    folder = packet_of_the_old_family(profile, packets, letter=name)
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert COVER_LETTER_COPY not in rendered(soffice)
    assert (folder / name).read_bytes() == b"THE OLD FAMILYS LETTER"
    assert not (folder / COVER_LETTER_COPY).exists()
    if letter == "txt":
        assert not (folder / LETTER_COPY_PDF).exists()


def test_only_the_postings_own_cover_letter_name_is_rendered_not_other_docx_files_in_the_packet(
    gdb, env, soffice, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb, role_family=NEW_FAMILY)
    folder = packet_with_letter(profile, packets, LETTER_COPY_DOCX, pdf=False)
    other_letter = COVER_LETTER_COPY.replace("Acme", "Other")
    (folder / "notes.docx").write_bytes(b"NOTES")
    (folder / other_letter).write_bytes(b"ANOTHER COMPANYS LETTER")
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert rendered(soffice).count(LETTER_COPY_DOCX) == 1
    assert "notes.docx" not in rendered(soffice)
    assert other_letter not in rendered(soffice)
    assert not (folder / "notes.pdf").exists()


def test_a_failed_render_of_an_earlier_role_familys_cover_letter_pdf_fails_the_row_and_leaves_no_partial_pdf(
    gdb, env, soffice, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb, role_family=NEW_FAMILY)
    folder = packet_with_letter(profile, packets, LETTER_COPY_DOCX, pdf=False)
    monkeypatch.setenv("SOFFICE_FAIL", "1")
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code != 0
    assert not (folder / LETTER_COPY_PDF).exists()
    assert (folder / LETTER_COPY_DOCX).read_bytes() == b"THE OLD FAMILYS LETTER"


def test_the_command_is_registered_with_id_and_dry_run(monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["jsa", "generate", "--help"])
    with pytest.raises(SystemExit) as stopped:
        cli.main()
    assert stopped.value.code == 0
    output = capsys.readouterr().out
    assert "--id" in output and "--dry-run" in output


# --- the redline (issues #82, #152; PRD 04 "Redline") -------------------------------

QUUX_JD = "# Staff Engineer\n\nOwn the Quuxlate system end to end.\n"
QUUX_EDIT = {
    "paragraph": 1,
    "find": "Quuxlate platform",
    "replace": "Quuxlate system",
    "goal": "ats",
    "jd_quote": "Quuxlate system",
    "reason": "REASON-MARKER",
}
QUUX_FLAG = {
    "paragraph": 1,
    "jd_quote": "Own the Quuxlate system end to end",
    "suggestion": "SUGGESTION-MARKER",
    "reason": "FLAG-REASON-MARKER",
}


def redline_reply(*edits, flags=(), explanation=None):
    return lambda prompt: result_message(
        result=json.dumps(
            {
                "edits": list(edits),
                "removals": [],
                "reorders": [],
                "flags": list(flags),
                "explanation": explanation,
            }
        )
    )


def proposed(folder):
    return json.loads((folder / REDLINE_EDITS).read_text(encoding="utf-8"))


def test_each_row_gets_one_redline_call_before_its_tracker_append(
    gdb, agent, gws, monkeypatch, capsys
):
    seed(gdb)
    log_at_redline = []

    def respond(prompt):
        log_at_redline.append(gws.exists())
        return result_message(result=REDLINE_JSON)

    agent.respond = respond
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert [call.options.model for call in agent.calls] == [REDLINE_MODEL]
    assert log_at_redline == [False]
    assert gws.exists()


def test_the_redline_run_uses_the_redline_settings_no_tools_one_turn_and_no_settings_files(
    gdb, agent, monkeypatch, capsys
):
    seed(gdb)
    jsa_generate(monkeypatch, capsys)
    (call,) = agent.calls
    assert call.options.model == REDLINE_MODEL
    assert call.options.effort == "medium"
    assert call.options.tools == []
    assert not call.options.allowed_tools
    assert call.options.max_turns == 1
    assert call.options.setting_sources == []


def test_the_redline_prompt_holds_the_posting_and_the_indexed_resume_paragraphs(
    gdb, agent, monkeypatch, capsys
):
    seed(gdb)
    jsa_generate(monkeypatch, capsys)
    (prompt,) = agent.prompts
    assert "Build the Frobnicator platform." in prompt
    assert "Riley Resumeperson" in prompt
    assert BASE_LINE in prompt
    assert "{{" not in prompt and "}}" not in prompt
    for private in ("Pat Example", "Jordan Example", "sheet-123"):
        assert private not in prompt


def test_every_proposed_edit_is_recorded_with_its_validation_result(
    gdb, env, agent, monkeypatch, capsys
):
    profile, packets = env
    seed(gdb)
    agent.respond = redline_reply(
        {**QUUX_EDIT, "paragraph": 99}, {**QUUX_EDIT, "reason": ""}
    )
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    record = proposed(packets / PLAIN)
    assert record["explanation"] is None
    assert len(record["edits"]) == 2
    assert all(edit["validation"] for edit in record["edits"])
    assert (packets / PLAIN / RESUME_COPY).read_bytes() == (
        family_dir(profile) / "resume.docx"
    ).read_bytes()


def test_a_run_with_no_valid_edit_writes_the_record_leaves_the_copy_and_renders_its_pdf(
    gdb, env, gws, soffice, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    folder = packets / PLAIN
    assert REDLINE_EDITS in entries(folder)
    assert proposed(folder)["explanation"] is None
    assert (folder / RESUME_COPY).read_bytes() == (
        family_dir(profile) / "resume.docx"
    ).read_bytes()
    assert (folder / RESUME_PDF).is_file()
    assert rendered(soffice) == [RESUME_COPY]
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_a_run_where_the_agent_proposes_no_edits_records_why_and_renders_the_resume_pdf(
    gdb, env, agent, gws, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb)
    agent.respond = redline_reply(explanation=NO_EDITS_EXPLANATION)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert proposed(packets / PLAIN) == {
        "explanation": NO_EDITS_EXPLANATION,
        "edits": [],
        "removals": [],
        "reorders": [],
        "flags": [],
    }
    folder = packets / PLAIN
    assert (folder / RESUME_COPY).read_bytes() == (
        family_dir(profile) / "resume.docx"
    ).read_bytes()
    assert (folder / RESUME_PDF).is_file()
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_a_valid_edit_becomes_a_tracked_change_with_a_comment_in_the_resume_copy(
    gdb, env, agent, soffice, monkeypatch, capsys
):
    _, packets = env
    seed(gdb, jd=QUUX_JD)
    agent.respond = redline_reply(QUUX_EDIT)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    folder = packets / PLAIN
    record = proposed(folder)
    assert record["explanation"] is None
    (edit,) = record["edits"]
    assert edit["validation"] is None
    with zipfile.ZipFile(folder / RESUME_COPY) as archive:
        body = archive.read("word/document.xml").decode("utf-8")
        comments = archive.read("word/comments.xml").decode("utf-8")
    assert "<w:ins " in body and "<w:del " in body
    assert "Claude (Redline)" in body
    assert "Quuxlate system" in comments and "REASON-MARKER" in comments
    # The packet holds one resume file, the redlined `.docx`, and no resume PDF.
    assert entries(folder) == {
        "job_posting.md",
        RESUME_COPY,
        REDLINE_EDITS,
    }
    assert rendered(soffice) == []


def test_a_flag_becomes_a_comment_without_a_change_and_the_resume_gets_no_pdf(
    gdb, env, agent, soffice, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb, jd=QUUX_JD)
    agent.respond = redline_reply(flags=[QUUX_FLAG])
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    folder = packets / PLAIN
    record = proposed(folder)
    assert record["explanation"] is None and record["edits"] == []
    (flag,) = record["flags"]
    assert flag["validation"] is None
    assert (folder / RESUME_COPY).read_bytes() != (
        family_dir(profile) / "resume.docx"
    ).read_bytes()
    with zipfile.ZipFile(folder / RESUME_COPY) as archive:
        body = archive.read("word/document.xml").decode("utf-8")
        comments = archive.read("word/comments.xml").decode("utf-8")
    assert "<w:ins " not in body and "<w:del " not in body
    assert "SUGGESTION-MARKER" in comments and "FLAG-REASON-MARKER" in comments
    assert entries(folder) == {"job_posting.md", RESUME_COPY, REDLINE_EDITS}
    assert rendered(soffice) == []
    assert re.search(r"\b1\b[^\n\d]*applied[^\n\d]*\b0\b[^\n\d]*dropped", output)
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_the_applied_and_dropped_counts_cover_edits_and_flags_together(
    gdb, agent, monkeypatch, capsys
):
    seed(gdb, jd=QUUX_JD)
    agent.respond = redline_reply(
        QUUX_EDIT,
        {**QUUX_EDIT, "paragraph": 99},
        flags=[QUUX_FLAG, {**QUUX_FLAG, "paragraph": 99}, {**QUUX_FLAG, "reason": ""}],
    )
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert re.search(r"\b2\b[^\n\d]*applied[^\n\d]*\b3\b[^\n\d]*dropped", output)


def test_a_run_where_every_proposal_is_dropped_leaves_the_copy_and_renders_its_pdf(
    gdb, env, agent, soffice, monkeypatch, capsys
):
    profile, packets = env
    seed(gdb, jd=QUUX_JD)
    agent.respond = redline_reply(
        {**QUUX_EDIT, "paragraph": 99}, flags=[{**QUUX_FLAG, "paragraph": 99}]
    )
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    folder = packets / PLAIN
    assert (folder / RESUME_COPY).read_bytes() == (
        family_dir(profile) / "resume.docx"
    ).read_bytes()
    assert (folder / RESUME_PDF).is_file()
    assert rendered(soffice) == [RESUME_COPY]


def test_the_redline_replaces_the_packets_resume_copy_and_never_the_profiles_base(
    gdb, env, agent, monkeypatch, capsys
):
    profile, packets = env
    base = (family_dir(profile) / "resume.docx").read_bytes()
    seed(gdb, jd=QUUX_JD)
    agent.respond = redline_reply(QUUX_EDIT)
    jsa_generate(monkeypatch, capsys)
    folder = packets / PLAIN
    assert (folder / RESUME_COPY).read_bytes() != base
    assert (family_dir(profile) / "resume.docx").read_bytes() == base
    assert not [name for name in entries(folder) if name.endswith("_redline.docx")]


@pytest.mark.parametrize(
    "text",
    [
        "not json at all",
        '{"paragraph": 1}',
        "",
        "   \n",
        "[]",
        json.dumps(
            {
                "edits": [],
                "removals": [],
                "reorders": [],
                "flags": [],
                "explanation": None,
            }
        ),
        json.dumps(
            {
                "edits": [],
                "removals": [],
                "reorders": [],
                "flags": [],
                "explanation": "",
            }
        ),
        json.dumps({"edits": [], "flags": [], "explanation": NO_EDITS_EXPLANATION}),
        json.dumps(
            {
                "edits": [],
                "reorders": [],
                "flags": [],
                "explanation": NO_EDITS_EXPLANATION,
            }
        ),
        json.dumps(
            {
                "edits": [],
                "removals": [],
                "flags": [],
                "explanation": NO_EDITS_EXPLANATION,
            }
        ),
        json.dumps({"edits": [], "explanation": NO_EDITS_EXPLANATION}),
        json.dumps(
            {"edits": [], "skills_edits": [], "explanation": NO_EDITS_EXPLANATION}
        ),
    ],
    ids=[
        "prose",
        "object",
        "empty",
        "blank",
        "bare-array",
        "no-edits-null-explanation",
        "no-edits-empty-explanation",
        "the-three-key-shape",
        "no-removals-key",
        "no-reorders-key",
        "no-flags-key",
        "the-previous-shape",
    ],
)
def test_a_redline_result_that_is_not_an_edits_and_explanation_object_flags_the_row_and_exits_non_zero(
    gdb, env, agent, gws, text, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    agent.respond = lambda prompt: result_message(result=text)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(posting_id) in output
    assert not gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 0
    assert REDLINE_EDITS not in entries(packets / PLAIN)


def test_a_failed_redline_run_flags_the_row_and_writes_no_record(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    agent.respond = lambda prompt: failed(api_error_status=529, result="overloaded")
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert "529" in output
    assert not gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 0
    assert REDLINE_EDITS not in entries(packets / PLAIN)


def test_a_failed_redline_on_one_row_does_not_stop_the_others(
    gdb, env, agent, gws, monkeypatch, capsys
):
    bad = seed(gdb, company="Bad Labs", jd="# Role\n\nBAD-JD-MARKER\n")
    good = seed(gdb, company="Good Labs")
    agent.respond = lambda prompt: (
        failed(result="overloaded")
        if "BAD-JD-MARKER" in prompt
        else result_message(result=REDLINE_JSON)
    )
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(bad) in output
    assert column(gdb, bad, "added_to_tracker") == 0
    assert column(gdb, good, "added_to_tracker") == 1


def test_a_rerun_after_a_failed_redline_resumes_at_the_redline(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    agent.respond = lambda prompt: result_message(result="nope")
    jsa_generate(monkeypatch, capsys)
    agent.respond = lambda prompt: result_message(result=REDLINE_JSON)
    agent.calls.clear()
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert len(agent.calls) == 1
    assert REDLINE_EDITS in entries(packets / PLAIN)
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_the_queue_path_skips_the_redline_when_the_record_exists(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    (folder / REDLINE_EDITS).write_text(
        NO_EDITS_RECORD,
        encoding="utf-8",
    )
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert agent.calls == []
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_an_id_refreshes_the_posting_but_makes_no_model_call_while_the_record_exists(
    gdb, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    jsa_generate(monkeypatch, capsys)
    (packets / PLAIN / "job_posting.md").write_text("stale text", encoding="utf-8")
    agent.calls.clear()
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert agent.calls == []
    assert (packets / PLAIN / "job_posting.md").read_text(encoding="utf-8") == JD


def test_an_id_reruns_the_redline_once_the_record_is_deleted(
    gdb, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    jsa_generate(monkeypatch, capsys)
    (packets / PLAIN / REDLINE_EDITS).unlink()
    agent.calls.clear()
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert len(agent.calls) == 1
    assert REDLINE_EDITS in entries(packets / PLAIN)


def applied_record() -> str:
    """A redline record in which one edit was applied."""
    return json.dumps(
        {
            "explanation": None,
            "edits": [{**QUUX_EDIT, "validation": None}],
            "removals": [],
            "reorders": [],
            "flags": [],
        }
    )


def test_an_id_leaves_a_redlined_copy_and_its_record_alone_but_refreshes_the_posting(
    gdb, env, agent, soffice, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    copy = folder / RESUME_COPY
    shutil.copy(family_dir(profile) / "resume.docx", copy)
    edit_copy(copy, "MY REVIEW IN PROGRESS")
    reviewed = copy.read_bytes()
    (folder / REDLINE_EDITS).write_text(applied_record(), encoding="utf-8")
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert agent.calls == []
    assert copy.read_bytes() == reviewed
    assert (folder / REDLINE_EDITS).read_text(encoding="utf-8") == applied_record()
    assert (folder / "job_posting.md").read_text(encoding="utf-8") == JD
    # An edit was applied, so the packet gets no resume PDF.
    assert RESUME_PDF not in entries(folder)
    assert rendered(soffice) == []


def test_a_record_whose_only_applied_proposal_is_a_flag_gets_no_resume_pdf(
    gdb, env, agent, soffice, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    shutil.copy(family_dir(profile) / "resume.docx", folder / RESUME_COPY)
    record = json.dumps(
        {
            "explanation": None,
            "edits": [],
            "removals": [],
            "reorders": [],
            "flags": [{**QUUX_FLAG, "validation": None}],
        }
    )
    (folder / REDLINE_EDITS).write_text(record, encoding="utf-8")
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert agent.calls == []
    assert RESUME_PDF not in entries(folder)
    assert rendered(soffice) == []
    assert (folder / REDLINE_EDITS).read_text(encoding="utf-8") == record


PREVIOUS_RECORDS = {
    "skills-edits": {
        "explanation": None,
        "edits": [],
        "skills_edits": [{**QUUX_EDIT, "validation": None}],
    },
    "three-keys": {
        "explanation": NO_EDITS_EXPLANATION,
        "edits": [],
        "flags": [],
    },
}


@pytest.mark.parametrize("args", ["queue", "id"])
@pytest.mark.parametrize("shape", sorted(PREVIOUS_RECORDS))
def test_a_record_in_the_previous_shape_is_not_read_and_flags_the_row(
    gdb, env, agent, gws, soffice, monkeypatch, capsys, args, shape
):
    profile, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    shutil.copy(family_dir(profile) / "resume.docx", folder / RESUME_COPY)
    record = json.dumps(PREVIOUS_RECORDS[shape])
    (folder / REDLINE_EDITS).write_text(record, encoding="utf-8")
    extra = ("--id", posting_id) if args == "id" else ()
    code, output = jsa_generate(monkeypatch, capsys, *extra)
    assert code != 0
    assert "is not a redline record" in output
    assert agent.calls == []
    assert rendered(soffice) == []
    assert (folder / REDLINE_EDITS).read_text(encoding="utf-8") == record
    assert column(gdb, posting_id, "added_to_tracker") == 0


@pytest.mark.parametrize("args", ["queue", "id"])
def test_a_record_with_no_applied_edit_and_no_resume_pdf_gets_the_pdf_without_a_redline(
    gdb, env, agent, soffice, monkeypatch, capsys, args
):
    profile, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    shutil.copy(family_dir(profile) / "resume.docx", folder / RESUME_COPY)
    record = NO_EDITS_RECORD
    (folder / REDLINE_EDITS).write_text(record, encoding="utf-8")
    extra = ("--id", posting_id) if args == "id" else ()
    code, _ = jsa_generate(monkeypatch, capsys, *extra)
    assert code == 0
    assert agent.calls == []
    assert (folder / RESUME_PDF).is_file()
    assert rendered(soffice) == [RESUME_COPY]
    assert (folder / REDLINE_EDITS).read_text(encoding="utf-8") == record


@pytest.mark.parametrize("args", ["queue", "id"])
def test_an_existing_resume_pdf_is_never_replaced_or_refreshed(
    gdb, env, agent, soffice, monkeypatch, capsys, args
):
    profile, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    shutil.copy(family_dir(profile) / "resume.docx", folder / RESUME_COPY)
    (folder / RESUME_PDF).write_bytes(b"MY OWN EXPORT")
    (folder / REDLINE_EDITS).write_text(
        NO_EDITS_RECORD,
        encoding="utf-8",
    )
    extra = ("--id", posting_id) if args == "id" else ()
    jsa_generate(monkeypatch, capsys, *extra)
    assert (folder / RESUME_PDF).read_bytes() == b"MY OWN EXPORT"
    assert rendered(soffice) == []


def test_the_redline_reads_the_resume_copy_as_it_stands_on_an_id(
    gdb, env, agent, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    jsa_generate(monkeypatch, capsys)
    edit_copy(packets / PLAIN / RESUME_COPY, "REVISED: led the Zorblax migration")
    (packets / PLAIN / REDLINE_EDITS).unlink()
    agent.calls.clear()
    jsa_generate(monkeypatch, capsys, "--id", posting_id)
    (prompt,) = agent.prompts
    assert "REVISED: led the Zorblax migration" in prompt


@pytest.mark.parametrize("failure", ["exits non-zero", "produces no PDF"])
def test_a_failed_resume_render_flags_the_row_leaves_no_partial_pdf_and_keeps_it_queued(
    gdb, env, gws, monkeypatch, capsys, failure
):
    _, packets = env
    posting_id = seed(gdb)
    monkeypatch.setenv(
        "SOFFICE_FAIL" if failure == "exits non-zero" else "SOFFICE_NO_PDF", "1"
    )
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(posting_id) in output
    assert not gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 0
    assert RESUME_PDF not in entries(packets / PLAIN)
    assert not list(packets.rglob(".soffice-*"))


def test_a_missing_soffice_flags_the_row_leaves_no_partial_pdf_and_keeps_it_queued(
    gdb, env, gws, tmp_path, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    monkeypatch.setenv("JSA_SOFFICE_BIN", str(tmp_path / "no-such-soffice"))
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(posting_id) in output
    assert not gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 0
    assert RESUME_PDF not in entries(packets / PLAIN)
    assert not list(packets.rglob(".soffice-*"))


def test_a_blind_row_gets_no_redline(gdb, env, agent, monkeypatch, capsys):
    _, packets = env
    seed(gdb, jd=None)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert agent.calls == []
    assert REDLINE_EDITS not in entries(packets / PLAIN)
    assert RESUME_PDF not in entries(packets / PLAIN)


def test_jsa_packet_writes_no_redline_or_resume_pdf_and_makes_no_model_call(
    gdb, env, agent, soffice, monkeypatch, capsys
):
    _, packets = env
    seed(gdb)
    jsa(monkeypatch, capsys, "packet")
    assert agent.calls == []
    assert entries(packets / PLAIN) == {"job_posting.md", RESUME_COPY}
    assert rendered(soffice) == []


def test_a_dry_run_makes_no_redline_call(gdb, env, agent, monkeypatch, capsys):
    _, packets = env
    seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys, "--dry-run")
    assert code == 0
    assert agent.calls == []
    assert entries(packets) == set()


def test_a_resume_copy_with_unresolved_tracked_changes_is_not_redlined_and_the_row_is_tracked(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    copy = folder / RESUME_COPY
    document = docx.Document()
    paragraph = document.add_paragraph("Built the ")
    tracked = parse_xml(
        '<w:ins xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' w:id="901" w:author="Someone" w:date="2026-01-01T00:00:00Z">'
        "<w:r><w:t>Quuxlate platform</w:t></w:r></w:ins>"
    )
    paragraph._p.append(tracked)
    document.save(copy)
    skipped = copy.read_bytes()
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert "warning" in output.lower()
    assert REDLINE_EDITS not in entries(folder)
    assert RESUME_PDF not in entries(folder)
    assert copy.read_bytes() == skipped
    assert column(gdb, posting_id, "added_to_tracker") == 1


# --- the optional wall-clock limit (issue #121; PRD 04 "Resume redline") ---


def limit_config(profile, packets, *, redline=""):
    text = CONFIG.format(packets=packets).replace(
        'effort = "medium"', f'effort = "medium"\n{redline}'
    )
    write_config_toml(profile, text)


def hang(monkeypatch, agent):
    """Every run never answers."""

    async def query(*, prompt, options=None, **_ignored):
        agent.calls.append(SimpleNamespace(prompt=prompt, options=options))
        await asyncio.sleep(3600)
        yield  # pragma: no cover

    monkeypatch.setattr(agent_loop, "query", query)


def test_a_timed_out_redline_does_not_stop_the_other_rows(
    gdb, env, agent, gws, monkeypatch, capsys
):
    profile, packets = env
    limit_config(profile, packets, redline="wall_clock_seconds = 0.2")
    slow = seed(gdb, company="Slow Labs", jd="# Role\n\nSLOW-JD-MARKER\n")
    fine = seed(gdb, company="Fine Labs")

    async def query(*, prompt, options=None, **ignored):
        if "SLOW-JD-MARKER" in prompt:
            await asyncio.sleep(3600)
        async for message in agent.query(prompt=prompt, options=options, **ignored):
            yield message

    monkeypatch.setattr(agent_loop, "query", query)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(slow) in output
    assert column(gdb, slow, "added_to_tracker") == 0
    assert column(gdb, fine, "added_to_tracker") == 1


def test_a_redline_that_runs_past_its_limit_flags_the_row_and_leaves_it_in_the_queue(
    gdb, env, agent, gws, monkeypatch, capsys
):
    profile, packets = env
    limit_config(profile, packets, redline="wall_clock_seconds = 0.05")
    posting_id = seed(gdb)
    hang(monkeypatch, agent)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code != 0
    assert str(posting_id) in output
    assert not gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 0
    assert REDLINE_EDITS not in entries(packets / PLAIN)


def test_a_rerun_after_a_timed_out_redline_resumes_at_the_redline(
    gdb, env, agent, gws, monkeypatch, capsys
):
    profile, packets = env
    limit_config(profile, packets, redline="wall_clock_seconds = 0.05")
    posting_id = seed(gdb)
    hang(monkeypatch, agent)
    jsa_generate(monkeypatch, capsys)
    monkeypatch.setattr(agent_loop, "query", agent.query)
    agent.calls.clear()
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert len(agent.calls) == 1
    assert REDLINE_EDITS in entries(packets / PLAIN)
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_a_run_inside_its_limit_completes_the_packet(
    gdb, env, agent, gws, monkeypatch, capsys
):
    profile, packets = env
    limit_config(profile, packets, redline="wall_clock_seconds = 60")
    posting_id = seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert REDLINE_EDITS in entries(packets / PLAIN)
    assert column(gdb, posting_id, "added_to_tracker") == 1


# --- placeholders (issue #130; PRD 04 "Placeholders in the resume and cover letter") ---


def save_with_company_token(path: Path) -> None:
    document = docx.Document()
    document.add_paragraph(BASE_LINE)
    document.add_paragraph("Keen on [COMPANY] and the [TITLE] role.")
    document.save(path)


def test_the_redline_reads_the_filled_copy(gdb, env, agent, monkeypatch, capsys):
    profile, packets = env
    save_with_company_token(family_dir(profile) / "resume.docx")
    seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    filled = "Keen on Acme Widgets and the Staff Engineer role."
    copy = docx.Document(str(packets / PLAIN / RESUME_COPY))
    assert copy.paragraphs[1].text == filled
    [prompt] = agent.prompts
    assert filled in prompt
    assert "[COMPANY]" not in prompt and "[TITLE]" not in prompt


def test_generate_fills_a_docx_cover_letter_copy_and_leaves_the_profiles_alone(
    gdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_company_token(family_dir(profile) / "cover_letter.docx")
    base = (family_dir(profile) / "cover_letter.docx").read_bytes()
    seed(gdb)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    letter = docx.Document(str(packets / PLAIN / COVER_LETTER_COPY))
    assert (
        letter.paragraphs[1].text == "Keen on Acme Widgets and the Staff Engineer role."
    )
    assert (family_dir(profile) / "cover_letter.docx").read_bytes() == base


def test_generate_never_fills_an_existing_copy_a_second_time(
    gdb, env, monkeypatch, capsys
):
    profile, packets = env
    save_with_company_token(family_dir(profile) / "resume.docx")
    save_with_company_token(family_dir(profile) / "cover_letter.docx")
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    kept = {}
    for name in (RESUME_COPY, COVER_LETTER_COPY):
        save_with_company_token(folder / name)
        kept[name] = (folder / name).read_bytes()
    for args in ((), ("--id", posting_id)):
        jsa_generate(monkeypatch, capsys, *args)
        for name, content in kept.items():
            assert (folder / name).read_bytes() == content
            assert "[COMPANY]" in docx.Document(str(folder / name)).paragraphs[1].text


# --- the skills section (issue #132; PRD 04 "Skills section") -----------------------

SKILLS_LINE = "Excel, Looker"
SK_EDIT = {
    "paragraph": 3,
    "find": "Looker",
    "replace": "Looker, Python",
    "goal": "narrative",
    "jd_quote": "Frobnicator platform",
    "reason": "GEN-SK-REASON-MARKER",
}
EXAMPLE_SKILLS = {
    "SQL",
    "dbt",
    "Looker",
    "Tableau",
    "Python",
    "A/B testing",
    "experimentation",
    "Data modeling",
    "Airflow",
    "BigQuery",
    "Snowflake",
    "Excel",
    "Stakeholder reporting",
}


def write_resume_with(profile: Path, *blocks: str) -> None:
    """The base resume: the name, the distinctive line, then these paragraphs."""
    document = docx.Document()
    document.add_paragraph().add_run("Riley Resumeperson").bold = True
    document.add_paragraph(BASE_LINE)
    for text in blocks:
        document.add_paragraph(text)
    document.save(family_dir(profile) / "resume.docx")


def skills_warnings(output: str) -> list[str]:
    return [
        line
        for line in output.splitlines()
        if "warning" in line.lower() and "skills" in line.lower()
    ]


def test_the_skills_list_in_the_profile_reaches_the_redline_prompt_one_item_per_line(
    gdb, agent, monkeypatch, capsys
):
    seed(gdb)
    jsa_generate(monkeypatch, capsys)
    (prompt,) = agent.prompts
    assert EXAMPLE_SKILLS <= set(prompt.splitlines())
    assert "{{" not in prompt and "}}" not in prompt


def test_a_skills_file_in_either_format_is_parsed_into_the_prompt(
    gdb, env, agent, monkeypatch, capsys
):
    profile, _ = env
    (profile / "resume_skills.md").write_text("Rust\n- Go, Zig\n* rust\n")
    seed(gdb)
    jsa_generate(monkeypatch, capsys)
    (prompt,) = agent.prompts
    lines = prompt.splitlines()
    assert {"Rust", "Go", "Zig"} <= set(lines)
    assert lines.count("Rust") == 1
    assert "Stakeholder reporting" not in prompt


@pytest.mark.parametrize(
    "content", [None, "", "  \n,\n -\n"], ids=["absent", "empty", "no-items"]
)
def test_without_a_skills_list_the_prompt_carries_none_and_a_skills_edit_is_dropped(
    gdb, env, agent, monkeypatch, capsys, content
):
    profile, packets = env
    write_resume_with(profile, "Skills", SKILLS_LINE)
    skills_file = profile / "resume_skills.md"
    if content is None:
        skills_file.unlink()
    else:
        skills_file.write_text(content)
    seed(gdb)
    agent.respond = redline_reply(SK_EDIT)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    (prompt,) = agent.prompts
    assert "Stakeholder reporting" not in prompt
    assert "{{" not in prompt and "}}" not in prompt
    folder = packets / PLAIN
    (entry,) = proposed(folder)["edits"]
    assert isinstance(entry["validation"], str) and entry["validation"]
    assert (folder / RESUME_COPY).read_bytes() == (
        family_dir(profile) / "resume.docx"
    ).read_bytes()
    assert (folder / RESUME_PDF).is_file()


def test_a_skills_edit_is_written_into_the_resume_copy_and_the_resume_gets_no_pdf(
    gdb, env, agent, soffice, monkeypatch, capsys
):
    profile, packets = env
    write_resume_with(profile, "Skills", SKILLS_LINE)
    posting_id = seed(gdb)
    agent.respond = redline_reply(SK_EDIT)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    folder = packets / PLAIN
    record = proposed(folder)
    assert record["explanation"] is None and record["flags"] == []
    (entry,) = record["edits"]
    assert entry["validation"] is None
    with zipfile.ZipFile(folder / RESUME_COPY) as archive:
        body = archive.read("word/document.xml").decode("utf-8")
        comments = archive.read("word/comments.xml").decode("utf-8")
    assert "<w:ins " in body and "<w:del " in body and "Claude (Redline)" in body
    assert "GEN-SK-REASON-MARKER" in comments
    assert RESUME_PDF not in entries(folder)
    assert rendered(soffice) == []
    assert re.search(r"\b1\b[^\n\d]*applied[^\n\d]*\b0\b[^\n\d]*dropped", output)
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_the_applied_and_dropped_counts_include_skills_section_edits(
    gdb, env, agent, monkeypatch, capsys
):
    profile, _ = env
    write_resume_with(profile, "Skills", SKILLS_LINE)
    seed(gdb)
    bad = {**SK_EDIT, "replace": "Looker, Rust"}
    agent.respond = redline_reply(SK_EDIT, bad)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert re.search(r"\b1\b[^\n\d]*applied[^\n\d]*\b1\b[^\n\d]*dropped", output)


def test_a_skills_edit_that_fails_validation_leaves_the_copy_and_renders_its_pdf(
    gdb, env, agent, monkeypatch, capsys
):
    profile, packets = env
    write_resume_with(profile, "Skills", SKILLS_LINE)
    seed(gdb)
    agent.respond = redline_reply({**SK_EDIT, "replace": "Looker, Rust"})
    jsa_generate(monkeypatch, capsys)
    folder = packets / PLAIN
    assert (folder / RESUME_COPY).read_bytes() == (
        family_dir(profile) / "resume.docx"
    ).read_bytes()
    assert (folder / RESUME_PDF).is_file()


def test_a_resume_with_no_skills_paragraph_warns_naming_the_row_and_still_builds_the_packet(
    gdb, agent, gws, monkeypatch, capsys
):
    posting_id = seed(gdb)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    warnings = skills_warnings(output)
    assert any(str(posting_id) in line for line in warnings)
    assert gws.exists()
    assert column(gdb, posting_id, "added_to_tracker") == 1


def test_a_resume_with_two_skills_paragraphs_warns(gdb, env, monkeypatch, capsys):
    profile, _ = env
    write_resume_with(profile, "Skills", SKILLS_LINE, "Skills:", "Python")
    posting_id = seed(gdb)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert any(str(posting_id) in line for line in skills_warnings(output))


def test_a_resume_with_one_skills_paragraph_draws_no_skills_warning(
    gdb, env, monkeypatch, capsys
):
    profile, _ = env
    write_resume_with(profile, "Skills", SKILLS_LINE)
    seed(gdb)
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    assert skills_warnings(output) == []


def test_a_resume_without_a_skills_section_holds_every_paragraph_to_the_ordinary_rules(
    gdb, env, agent, monkeypatch, capsys
):
    _, packets = env
    seed(gdb, jd=QUUX_JD)
    removal = {
        **QUUX_EDIT,
        "find": " from scratch",
        "replace": "",
    }
    agent.respond = redline_reply(QUUX_EDIT, removal)
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    record = proposed(packets / PLAIN)
    assert record["edits"][0]["validation"] is None
    assert record["edits"][1]["validation"]


# --- removals and reorders in a packet (issue #153; PRD 04 "Redline", "Written into the resume copy") ---

BULLET_LINES = [
    "Built the Quuxlate platform from scratch.",
    "Led data work using SQL daily.",
    "Reduced costs by 30% in 2020.",
]


def write_bulleted_resume(profile: Path) -> None:
    document = docx.Document()
    document.add_paragraph().add_run("Riley Resumeperson").bold = True
    for text in BULLET_LINES:
        document.add_paragraph(text, style="List Bullet")
    document.save(family_dir(profile) / "resume.docx")


def structural_reply(removals=(), reorders=()):
    return lambda prompt: result_message(
        result=json.dumps(
            {
                "edits": [],
                "removals": list(removals),
                "reorders": list(reorders),
                "flags": [],
                "explanation": None,
            }
        )
    )


def test_an_applied_removal_and_reorder_give_a_lone_docx_and_count_in_the_report(
    gdb, env, agent, soffice, monkeypatch, capsys
):
    profile, packets = env
    write_bulleted_resume(profile)
    seed(gdb, jd=QUUX_JD)
    quote = "Own the Quuxlate system end to end"
    agent.respond = structural_reply(
        removals=[
            {"paragraph": 3, "jd_quote": quote, "reason": "REMOVAL-MARKER"},
            {"paragraph": 0, "jd_quote": quote, "reason": "the name is no bullet"},
        ],
        reorders=[
            {"paragraphs": [2, 1], "jd_quote": quote, "reason": "REORDER-MARKER"}
        ],
    )
    code, output = jsa_generate(monkeypatch, capsys)
    assert code == 0
    folder = packets / PLAIN
    record = proposed(folder)
    assert [r["validation"] is None for r in record["removals"]] == [True, False]
    assert [r["validation"] is None for r in record["reorders"]] == [True]
    with zipfile.ZipFile(folder / RESUME_COPY) as archive:
        body = archive.read("word/document.xml").decode("utf-8")
        comments = archive.read("word/comments.xml").decode("utf-8")
    assert "<w:moveFrom" not in body and "<w:moveTo" not in body
    assert "<w:del " in body
    assert body.index("Led data work") < body.index("Built the Quuxlate")
    assert "Claude (Redline)" in body
    assert "REMOVAL-MARKER" in comments and "REORDER-MARKER" in comments
    assert entries(folder) == {"job_posting.md", RESUME_COPY, REDLINE_EDITS}
    assert rendered(soffice) == []
    assert re.search(r"\b2\b[^\n\d]*applied[^\n\d]*\b1\b[^\n\d]*dropped", output)


def test_a_record_whose_only_applied_proposal_is_a_removal_gets_no_resume_pdf(
    gdb, env, agent, soffice, monkeypatch, capsys
):
    profile, packets = env
    posting_id = seed(gdb)
    folder = packets / PLAIN
    folder.mkdir(parents=True)
    shutil.copy(family_dir(profile) / "resume.docx", folder / RESUME_COPY)
    record = json.dumps(
        {
            "explanation": None,
            "edits": [],
            "removals": [
                {
                    "paragraph": 1,
                    "jd_quote": "x",
                    "reason": "y",
                    "validation": None,
                }
            ],
            "reorders": [],
            "flags": [],
        }
    )
    (folder / REDLINE_EDITS).write_text(record, encoding="utf-8")
    code, _ = jsa_generate(monkeypatch, capsys, "--id", posting_id)
    assert code == 0
    assert agent.calls == []
    assert RESUME_PDF not in entries(folder)
    assert rendered(soffice) == []


def test_a_posting_kept_out_of_refinement_gets_the_same_packet_and_tracker_row(
    gdb, env, agent, gws, monkeypatch, capsys
):
    _, packets = env
    ordinary = seed(gdb, company="Beta Labs")
    excluded = seed(gdb, company="Gamma Labs")
    gdb.execute("UPDATE postings SET in_refinement = 0 WHERE id = ?", (excluded,))
    code, _ = jsa_generate(monkeypatch, capsys)
    assert code == 0
    # The resume copy's name carries the company; every other entry is identical.
    assert entries(packets / "Beta Labs - Staff Engineer") == {
        name.replace("GammaLabs", "BetaLabs")
        for name in entries(packets / "Gamma Labs - Staff Engineer")
    }
    assert entries(packets / "Gamma Labs - Staff Engineer") == {
        "job_posting.md",
        "PatExample_Resume_StaffEngineer_GammaLabs.docx",
        "PatExample_Resume_StaffEngineer_GammaLabs.pdf",
        REDLINE_EDITS,
    }
    assert sorted(row[0] for row in appended_rows(gws)) == [ordinary, excluded]
    assert column(gdb, excluded, "added_to_tracker") == 1
    assert len(agent.calls) == 2
    assert column(gdb, excluded, "in_refinement") == 0


def test_packet_and_track_treat_a_posting_kept_out_of_refinement_like_any_other(
    gdb, env, gws, monkeypatch, capsys
):
    _, packets = env
    excluded = seed(gdb, company="Gamma Labs")
    gdb.execute("UPDATE postings SET in_refinement = 0 WHERE id = ?", (excluded,))
    code, _ = jsa(monkeypatch, capsys, "packet")
    assert code == 0
    assert (packets / "Gamma Labs - Staff Engineer" / "job_posting.md").exists()
    code, _ = jsa(monkeypatch, capsys, "track")
    assert code == 0
    assert [row[0] for row in appended_rows(gws)] == [excluded]
    assert column(gdb, excluded, "added_to_tracker") == 1
